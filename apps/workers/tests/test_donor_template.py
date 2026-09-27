"""Tests for donor-template rendering (docxtpl) — placeholder insertion,
rendering, and the FastAPI routes. Runs without an LLM."""
from __future__ import annotations

import base64
from io import BytesIO

from docx import Document

from app.donor_template.models import RegionRef
from app.donor_template.placeholder_inserter import insert_placeholders
from app.donor_template.renderer import render

AUTH = {"x-internal-token": "test-internal-token"}


def _fixture_docx() -> bytes:
    doc = Document()
    doc.add_paragraph("Donor Report Template", style="Title")
    doc.add_paragraph("1. Executive Summary", style="Heading 1")
    doc.add_paragraph("placeholder body text")
    doc.add_paragraph("2. Indicator Performance", style="Heading 1")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Code"
    table.rows[0].cells[1].text = "Value"
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# placeholder_inserter — pure, no FastAPI
# --------------------------------------------------------------------------- #


def test_insert_placeholders_inserts_tag_before_mapped_heading() -> None:
    original = _fixture_docx()
    regions = [RegionRef(id="h-0001", kind="HEADING", order=1, placeholderKey="region_h_0001")]
    templated = insert_placeholders(original, regions)
    doc = Document(BytesIO(templated))
    texts = [p.text for p in doc.paragraphs]
    assert "{{ region_h_0001 }}" in texts
    idx = texts.index("{{ region_h_0001 }}")
    assert texts[idx + 1] == "1. Executive Summary"


def test_insert_placeholders_inserts_tag_before_mapped_table() -> None:
    original = _fixture_docx()
    regions = [RegionRef(id="t-0000", kind="TABLE", order=2, placeholderKey="region_t_0000")]
    templated = insert_placeholders(original, regions)
    doc = Document(BytesIO(templated))
    assert len(doc.tables) == 1, "the table itself must not be duplicated or removed"
    texts = [p.text for p in doc.paragraphs]
    assert "{{ region_t_0000 }}" in texts


def test_insert_placeholders_ignores_regions_not_supplied() -> None:
    original = _fixture_docx()
    templated = insert_placeholders(original, [])
    doc = Document(BytesIO(templated))
    texts = [p.text for p in doc.paragraphs]
    assert not any(t.startswith("{{") for t in texts), "no placeholder should be inserted when no region is mapped"


def test_insert_placeholders_uses_title_style_as_heading_index_zero() -> None:
    """The TS mammoth parser maps Word's Title style to h1 in the same
    position; the Python side must count it the same way so region ids
    computed by one side locate the right block via the other."""
    original = _fixture_docx()
    regions = [RegionRef(id="h-0000", kind="HEADING", order=0, placeholderKey="region_h_0000")]
    templated = insert_placeholders(original, regions)
    doc = Document(BytesIO(templated))
    texts = [p.text for p in doc.paragraphs]
    idx = texts.index("{{ region_h_0000 }}")
    assert texts[idx + 1] == "Donor Report Template"


# --------------------------------------------------------------------------- #
# renderer — pure, no FastAPI
# --------------------------------------------------------------------------- #


def test_render_substitutes_placeholder_with_content() -> None:
    original = _fixture_docx()
    templated = insert_placeholders(original, [RegionRef(id="h-0001", kind="HEADING", order=1, placeholderKey="region_h_0001")])
    rendered = render(templated, {"region_h_0001": "Real generated content for this section."})
    doc = Document(BytesIO(rendered))
    texts = [p.text for p in doc.paragraphs]
    assert "Real generated content for this section." in texts
    assert not any("{{" in t for t in texts), "no unrendered Jinja2 tags should remain"


# --------------------------------------------------------------------------- #
# HTTP routes
# --------------------------------------------------------------------------- #


def test_insert_placeholders_route_requires_token(client) -> None:
    res = client.post("/v1/donor-template/insert-placeholders", json={"originalDocxBase64": "", "regions": []})
    assert res.status_code == 401


def test_insert_placeholders_route_end_to_end(client) -> None:
    original = _fixture_docx()
    res = client.post(
        "/v1/donor-template/insert-placeholders",
        headers=AUTH,
        json={
            "originalDocxBase64": base64.b64encode(original).decode("ascii"),
            "regions": [{"id": "h-0001", "kind": "HEADING", "order": 1, "placeholderKey": "region_h_0001"}],
        },
    )
    assert res.status_code == 200
    templated_b64 = res.json()["templatedDocxBase64"]
    doc = Document(BytesIO(base64.b64decode(templated_b64)))
    assert "{{ region_h_0001 }}" in [p.text for p in doc.paragraphs]


def test_render_route_end_to_end(client) -> None:
    original = _fixture_docx()
    templated = insert_placeholders(original, [RegionRef(id="h-0001", kind="HEADING", order=1, placeholderKey="region_h_0001")])
    res = client.post(
        "/v1/donor-template/render",
        headers=AUTH,
        json={
            "templatedDocxBase64": base64.b64encode(templated).decode("ascii"),
            "context": {"region_h_0001": "Rendered via the HTTP route."},
        },
    )
    assert res.status_code == 200
    rendered_b64 = res.json()["renderedDocxBase64"]
    doc = Document(BytesIO(base64.b64decode(rendered_b64)))
    assert "Rendered via the HTTP route." in [p.text for p in doc.paragraphs]


def test_render_route_returns_422_on_malformed_input(client) -> None:
    res = client.post(
        "/v1/donor-template/render",
        headers=AUTH,
        json={"templatedDocxBase64": "not-valid-base64-docx", "context": {}},
    )
    assert res.status_code == 422


# --------------------------------------------------------------------------- #
# Markdown → native Word blocks (report editor rich text)
# --------------------------------------------------------------------------- #

from app.donor_template.markdown_docx import parse_inline, parse_markdown_blocks  # noqa: E402

RICH_SECTION = """Households reached **1,680** of the *2,000* target.

### Key results

- Boreholes rehabilitated
  - Four in District A
1. Training delivered
> Community feedback was positive.

| Indicator | Target | Achieved |
|---|---|---|
| OUT-1 Households | 2,000 | 1,680 |
| OUT-2 Latrines & sanitation | 400 | 248 |

See [the log](https://example.org/log) for details <draft>."""


def test_parse_markdown_blocks_covers_the_editor_subset() -> None:
    kinds = [b.kind for b in parse_markdown_blocks(RICH_SECTION)]
    assert kinds == ["paragraph", "heading", "bullet", "bullet", "bullet", "quote", "table", "paragraph"]
    blocks = parse_markdown_blocks(RICH_SECTION)
    assert blocks[3].level == 1 and not blocks[3].ordered
    assert blocks[4].ordered and blocks[4].marker == "1."
    assert blocks[6].header == ["Indicator", "Target", "Achieved"]
    assert blocks[6].rows[1] == ["OUT-2 Latrines & sanitation", "400", "248"]


def test_parse_inline_emphasis_code_and_links() -> None:
    runs = parse_inline("A **bold** and *italic* `code` [label](http://x) end")
    assert [(r.text, r.bold, r.italic, r.code) for r in runs] == [
        ("A ", False, False, False),
        ("bold", True, False, False),
        (" and ", False, False, False),
        ("italic", False, True, False),
        (" ", False, False, False),
        ("code", False, False, True),
        (" ", False, False, False),
        ("label (http://x)", False, False, False),
        (" end", False, False, False),
    ]
    assert parse_inline("2 * 3 = 6") == [parse_inline("2 * 3 = 6")[0]]


def test_parse_is_total_on_odd_input() -> None:
    assert parse_markdown_blocks("") == []
    assert [b.kind for b in parse_markdown_blocks("| not a table\n**unclosed")] == ["paragraph"]


def _render_rich(content: str) -> Document:
    templated = insert_placeholders(_fixture_docx(), [RegionRef(id="h-0001", kind="HEADING", order=1, placeholderKey="region_h_0001")])
    return Document(BytesIO(render(templated, {"region_h_0001": content})))


def test_render_turns_markdown_into_native_blocks() -> None:
    doc = _render_rich(RICH_SECTION)
    texts = [p.text for p in doc.paragraphs]
    body = "\n".join(texts)
    assert "**" not in body and "|" not in body and "{{" not in body
    assert "Households reached 1,680 of the 2,000 target." in texts
    bold_runs = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert "1,680" in bold_runs
    heading = next(p for p in doc.paragraphs if p.text == "Key results")
    assert heading.style.name == "Heading 3"
    bullet = next(p for p in doc.paragraphs if p.text == "Boreholes rehabilitated")
    assert bullet.style.name == "List Bullet"
    assert any(p.text == "Four in District A" for p in doc.paragraphs)
    # Escaping: text with & and < survives as literal characters.
    assert "See the log (https://example.org/log) for details <draft>." in texts
    # The markdown table became a real Word table (the fixture already has one).
    rich = [t for t in doc.tables if t.rows[0].cells[0].text == "Indicator"]
    assert len(rich) == 1
    assert [c.text for c in rich[0].rows[2].cells] == ["OUT-2 Latrines & sanitation", "400", "248"]
    assert rich[0].rows[0].cells[0].paragraphs[0].runs[0].bold is True


def test_render_places_blocks_where_the_placeholder_was() -> None:
    doc = _render_rich("First paragraph.\n\nSecond paragraph.")
    texts = [p.text for p in doc.paragraphs]
    first = texts.index("First paragraph.")
    assert texts[first + 1] == "Second paragraph."
    # insert_placeholders puts the tag just before the mapped heading
    # (h-0001 = "1. Executive Summary"; the Title is h-0000).
    assert texts.index("1. Executive Summary") == first + 2


def test_render_falls_back_when_template_lacks_list_styles() -> None:
    templated = insert_placeholders(_fixture_docx(), [RegionRef(id="h-0001", kind="HEADING", order=1, placeholderKey="region_h_0001")])
    tpl_doc = Document(BytesIO(templated))
    # Simulate a donor template without list styles by renaming them.
    for name in ("List Bullet", "List Number", "List Bullet 2"):
        try:
            tpl_doc.styles[name].name = f"Donor {name}"
        except KeyError:
            pass
    buf = BytesIO()
    tpl_doc.save(buf)
    rendered = Document(BytesIO(render(buf.getvalue(), {"region_h_0001": "- One\n1. Two"})))
    texts = [p.text for p in rendered.paragraphs]
    assert "• One" in texts and "1. Two" in texts


def test_inline_placeholder_gets_plain_text() -> None:
    doc = Document()
    doc.add_paragraph("Summary: {{ region_x }} (auto)")
    buf = BytesIO()
    doc.save(buf)
    rendered = Document(BytesIO(render(buf.getvalue(), {"region_x": "**Bold** & more"})))
    assert rendered.paragraphs[0].text == "Summary: Bold & more (auto)"


def test_parse_inline_keeps_spaced_asterisks_literal() -> None:
    assert [(r.text, r.italic) for r in parse_inline("2 * 3 * 4 sessions")] == [("2 * 3 * 4 sessions", False)]
    assert [(r.text, r.bold) for r in parse_inline("**1,680** done")] == [("1,680", True), (" done", False)]
