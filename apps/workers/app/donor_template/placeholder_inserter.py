"""Physically inserts docxtpl `{{ placeholder }}` Jinja2 tags into a donor's
original DOCX at the locations the TS-side structural parser (mammoth-based)
detected as regions. Uses python-docx to walk the document body in true
document order, counting headings and tables with the SAME per-kind
numbering scheme the TS parser uses (`h-0000, h-0001, ...` / `t-0000,
t-0001, ...`), so a region id computed by one library's structural view
locates the same physical paragraph/table via the other's.

This module never re-detects or re-scores regions — it only relocates
regions the caller has already decided are the right ones (`RegionRef`).

python-docx ships no type stubs; internal helpers here are typed `Any` for
its objects rather than fighting mypy over an untyped third-party API
(consistent with `app/parsers.py`'s existing docx usage in this codebase).
"""
from __future__ import annotations

from io import BytesIO
from typing import Any, Iterator

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .models import RegionRef

_HEADING_STYLE_PREFIXES = ("heading", "title")


def _is_heading_paragraph(paragraph: Any) -> bool:
    style = paragraph.style
    style_name = ((style.name if style is not None else "") or "").strip().lower()
    return any(style_name.startswith(p) for p in _HEADING_STYLE_PREFIXES)


def _iter_body_blocks(document: Any) -> Iterator[tuple[str, Any]]:
    """Yields (kind, element) for every top-level paragraph/table in the
    document body, in true document order — python-docx's own
    `document.paragraphs`/`document.tables` are separate flat lists that do
    NOT preserve interleaving, so this walks the raw XML body instead."""
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield "p", Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield "tbl", Table(child, document)


def insert_placeholders(original_docx_bytes: bytes, regions: list[RegionRef]) -> bytes:
    document: Any = Document(BytesIO(original_docx_bytes))

    by_id = {r.id: r for r in regions}
    heading_index = 0
    table_index = 0

    for kind, element in _iter_body_blocks(document):
        if kind == "p":
            paragraph = element
            if not _is_heading_paragraph(paragraph):
                continue
            region_id = f"h-{heading_index:04d}"
            heading_index += 1
            region = by_id.get(region_id)
            if region is None or region.kind != "HEADING":
                continue
            new_paragraph = paragraph.insert_paragraph_before("")
            new_paragraph.add_run(f"{{{{ {region.placeholderKey} }}}}")
        else:
            table = element
            region_id = f"t-{table_index:04d}"
            table_index += 1
            region = by_id.get(region_id)
            if region is None or region.kind != "TABLE":
                continue
            # Insert the placeholder as a new paragraph immediately before
            # the table (docxtpl can substitute a whole table via a
            # `{%tr %}`-style row loop, but a single placeholder paragraph
            # covering "this table's content" is the v1 scope — matches
            # what `buildSectionSpecificGuidance` already asks the narrator
            # to produce: one markdown table per relevant section, rendered
            # as plain text into the placeholder).
            new_paragraph = document.add_paragraph("")
            new_paragraph.add_run(f"{{{{ {region.placeholderKey} }}}}")
            # `add_paragraph` appends at the end of the body; lxml's
            # `addprevious` detaches an element from its current parent and
            # relocates it, so this moves (not duplicates) the new
            # paragraph to immediately before the table.
            table._tbl.addprevious(new_paragraph._p)

    buf = BytesIO()
    document.save(buf)
    return buf.getvalue()
