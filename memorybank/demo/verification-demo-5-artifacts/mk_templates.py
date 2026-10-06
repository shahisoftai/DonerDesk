from docx import Document
def mk(path, title, intro, sections):
    d = Document(); d.add_heading(title, 0); d.add_paragraph(intro)
    for lvl, h, ins, qs in sections:
        d.add_heading(h, lvl); d.add_paragraph(ins)
        for q in qs: d.add_paragraph(q, style="List Bullet")
    d.save(path)
mk("USAID-monthly-progress-report-template.docx", "USAID/Kenya — Monthly Activity Progress Report",
 "Implementing partners report monthly against the approved Activity Work Plan and Performance Management Plan (PMP). Report only verified figures. Disaggregate people-level results by sex and age where the indicator requires it. Submit to the Agreement Officer's Representative (AOR) within 10 days of month end.",
 [(1, "Executive Summary", "Summarise the month's main achievements, challenges and priorities for next month in no more than 250 words. Acknowledge USAID funding.", []),
  (1, "Progress Against the Work Plan", "Report progress by Intermediate Result (IR) and activity against the Work Plan.", ["Narrative for each IR: what was delivered, where, and with whom.", "Provide an indicator performance table: indicator, month result, cumulative result, life-of-activity target, percent of target.", "Explain any indicator that is behind plan and the corrective action."]),
  (1, "Gender Equality and Women's Empowerment", "Describe how gender was integrated into activities this month.", ["Report sex-disaggregated results.", "What actions addressed gender gaps in access to services?"]),
  (1, "Environmental Compliance", "Report status of environmental mitigation measures under the Initial Environmental Examination (IEE).", ["Health care waste management (sharps, placenta pits, incinerators).", "Any environmental incidents or unexpected impacts."]),
  (1, "Collaboration, Learning and Adaptation", "Describe how the Activity collaborated with others, what it learned and how it adapted.", ["Coordination with the County Department of Health and other partners.", "What evidence led to a change in approach?"]),
  (1, "Implementation Challenges and Actions Taken", "Describe challenges, risks that materialised, and the action taken.", ["Security, access, supply chain and workforce issues.", "Any request for USAID support."]),
  (1, "Priorities for Next Month", "Set out planned activities and expected results for the coming month.", []),
  (1, "Branding and Marking Compliance", "Confirm that USAID branding and marking requirements were met on materials, equipment and events.", [])])
mk("USAID-activity-completion-report-template.docx", "USAID/Kenya — Activity Completion Report",
 "The Activity Completion Report (ACR) is submitted at the end of the award. It covers the whole life of the Activity: final performance against every PMP indicator, key achievements, lessons and recommendations. Submit within 30 days of the end date.",
 [(1, "Executive Summary", "Summarise the Activity, its main results and overall performance in no more than 400 words.", []),
  (1, "Background and Context", "Describe the health context in Turkana County, the needs addressed and the Activity design.", ["Where and with whom did the Activity work?"]),
  (1, "Final Performance Against Indicators", "Report life-of-activity results against each PMP indicator and its target.", ["Provide an indicator table: baseline, target, achieved, percent of target.", "Disaggregate people-level results by sex.", "Explain any target that was not met."]),
  (1, "Key Achievements by Intermediate Result", "Summarise achievements under each Intermediate Result.", []),
  (2, "Intermediate Result 1: Quality facility-based MNCH services", "Report on clinical training and mentorship, equipment, antenatal care, skilled delivery, immunization and facility data quality.", ["How many health workers were trained and mentored?", "How did service use change at supported facilities?"]),
  (2, "Intermediate Result 2: Community health system and household practices", "Report on community health volunteers, health education and malnutrition screening and referral.", ["How many CHVs were trained and how many people were reached?"]),
  (1, "Gender Equality and Social Inclusion", "Describe results for women and girls, and for people with disabilities.", ["Provide sex-disaggregated data and explain gender gaps."]),
  (1, "Environmental Compliance", "Report final status of environmental mitigation under the IEE.", []),
  (1, "Collaboration, Learning and Adaptation", "Describe collaboration, key learning and adaptations during the Activity.", ["What would you do differently?"]),
  (1, "Sustainability and Transition", "Explain how results will be sustained after the Activity ends and what was handed over to the County Government.", ["Which assets, systems and responsibilities were transferred?"]),
  (1, "Challenges and Lessons Learned", "Describe the main challenges, how they were addressed and lessons for future programming.", []),
  (1, "Recommendations for Future Programming", "Recommend priorities for USAID and the County Government.", []),
  (1, "Financial Summary", "Provide a summary of funds obligated, expended and unspent, and cost share (SF-425 style).", ["Include a table: budget line, budget, expended, balance.", "Explain any variance above ten percent."]),
  (1, "Branding and Marking Compliance", "Confirm that USAID branding and marking requirements were met throughout the Activity.", [])])
