from docx import Document
def mk(path, title, intro, sections):
    d = Document(); d.add_heading(title, 0); d.add_paragraph(intro)
    for lvl, h, ins, qs in sections:
        d.add_heading(h, lvl); d.add_paragraph(ins)
        for q in qs: d.add_paragraph(q, style="List Bullet")
    d.save(path)
mk("EU-monthly-progress-report-template.docx", "European Union — Monthly Progress Report (Grant EU-GH/2025/HD/0417)",
 "The Beneficiary reports monthly to the Contracting Authority on progress of the Action against the Logical Framework Matrix. Report only verified figures. Use the formal register of EU reporting: refer to 'the Action', write in the third person and avoid first person. Disaggregate people-level results by sex. Submit within 10 days of month end.",
 [(1,"Summary of the Month","Summarise the month's main results, deviations and priorities in no more than 250 words. Acknowledge the European Union as donor.",["Which results were achieved against the monthly plan, and where?"]),
  (1,"Progress Against Expected Results","Report progress by Outcome and Output of the Logical Framework Matrix.",["Provide a results table: indicator, this month, cumulative, target, percent of target.","Explain any indicator behind plan and the corrective action."]),
  (1,"Activities Implemented","List the activities carried out in the month, by district, with dates and participants.",["Which activities were planned but not carried out, and why?"]),
  (1,"Outcome Indicators and Data Quality","Report the outcome-level indicators from DHIMS2 and surveys and comment on data completeness and reliability.",["Which indicators were not measured this month, and why?"]),
  (1,"Challenges, Risks and Mitigation","Describe challenges and risks that materialised and the mitigation applied.",["Any request for support or a change that needs the Contracting Authority's agreement?"]),
  (1,"Cross-Cutting Issues: Gender, Environment and Do-No-Harm","Describe how gender equality, climate and environment (health-care waste) and the do-no-harm principle were integrated this month.",["Provide sex-disaggregated results.","Any incident or safeguarding concern?"]),
  (1,"Visibility and Communication","Confirm that EU visibility requirements were met on materials, equipment and events.",["Which communication activities took place?"]),
  (1,"Plan for the Next Month","Set out planned activities and expected results for the coming month.",[])])
mk("EU-final-narrative-report-template.docx", "European Union — Final Narrative Report (Grant EU-GH/2025/HD/0417)",
 "The Final Narrative Report covers the whole duration of the Action and is submitted within 3 months of its end date. It assesses results against the Logical Framework Matrix and the OECD-DAC evaluation criteria. Use the formal register of EU reporting: 'the Action', third person, evidence-based statements only. Report life-of-Action figures from verified records.",
 [(1,"Executive Summary","Summarise the Action, its main results and overall performance in no more than 400 words.",[]),
  (1,"Context and Description of the Action","Describe the context and the design of the Action.",[]),
  (2,"Context Update","Describe the health context in Northern Region and any significant change during implementation.",["Where and with whom did the Action work?"]),
  (2,"Intervention Logic","Describe how the Action's outputs were expected to lead to outcomes and impact, and whether the logic held.",[]),
  (1,"Assessment of Implementation","Assess implementation against the Logical Framework Matrix.",[]),
  (2,"Impact and Outcome Results","Report life-of-Action results for the impact and outcome indicators.",["Provide an indicator table: baseline, target, achieved, percent of target.","Explain any target that was not met."]),
  (2,"Output Delivery","Report life-of-Action delivery of every output, disaggregated by sex where applicable.",["Provide a table of output indicators: target, achieved, percent of target."]),
  (2,"Activities","Summarise the activities implemented by Outcome.",[]),
  (2,"Deviations from Plan and Variance","Explain deviations from the planned work plan and any variance in results.",["What caused each deviation and what was done about it?"]),
  (1,"Assessment Against Evaluation Criteria","Assess the Action against the OECD-DAC criteria.",[]),
  (2,"Relevance","Was the Action relevant to the needs of women and children and to national policy?",[]),
  (2,"Effectiveness","To what extent were the expected outcomes achieved?",[]),
  (2,"Efficiency","Were results delivered at reasonable cost and on time?",[]),
  (2,"Sustainability","Which results are likely to continue after the Action ends?",[]),
  (2,"Impact","What changes in maternal and child health can be attributed to the Action?",[]),
  (1,"Cross-Cutting Issues","Report on gender equality, climate and environment, and the rights-based approach.",["Provide sex-disaggregated results and explain gender gaps."]),
  (1,"Risks and Assumptions","Report on the risks and assumptions of the Logical Framework Matrix and which materialised.",[]),
  (1,"Visibility and Communication","Confirm that EU visibility requirements were met throughout the Action.",[]),
  (1,"Lessons Learned","Describe the main lessons for future programming.",["What would you do differently?"]),
  (1,"Sustainability and Exit Strategy","Explain what was handed over to the District Health Management Teams and the Ghana Health Service, and how results will be sustained.",["Which assets, systems and responsibilities were transferred?"]),
  (1,"Financial Summary","Provide a summary of the budget and expenditure of the Action.",["Include a table: budget line, budget, expended, balance.","Explain any variance above ten percent."]),
  (1,"Recommendations","Recommend priorities for the Contracting Authority, the Ghana Health Service and implementing partners.",[])])
