from docx import Document
def mk(path, title, sections, intro):
    d = Document(); d.add_heading(title, 0); d.add_paragraph(intro)
    for h, ins, qs in sections:
        d.add_heading(h, 1); d.add_paragraph(ins)
        for q in qs: d.add_paragraph(q, style="List Bullet")
    d.save(path)
mk("GWHF-monthly-progress-template.docx", "Global Water & Health Fund — Monthly Progress Report",
 [("Executive Summary", "Summarise the month's main achievements, challenges and priorities for next month in no more than 200 words.", []),
  ("Progress Against Results", "Report progress against each logframe indicator for the month and cumulatively against target.", ["Provide an indicator table: indicator, month value, cumulative, target, % of target.", "Explain any indicator behind plan."]),
  ("Activities Implemented", "Describe the activities delivered during the month, where, and with whom.", ["Which water, sanitation and hygiene activities were completed?", "How many people participated, by sex?"]),
  ("Challenges and Mitigation", "Describe challenges encountered and the actions taken.", ["What risks materialised?", "What was done about them?"]),
  ("Lessons Learned", "Note lessons learned during the month that could improve delivery.", []),
  ("Next Month's Plan", "Set out planned activities for the coming month.", [])],
 "Submit within 10 days of month end. Use plain language and report only verified figures.")
mk("GWHF-final-report-template.docx", "Global Water & Health Fund — Final Project Report",
 [("Executive Summary", "Summarise the project, its main results and overall performance in no more than 300 words.", []),
  ("Project Background and Context", "Describe the context, the needs addressed and the project design.", ["Where and with whom did the project work?"]),
  ("Results Achieved", "Report life-of-project results against each logframe indicator and its target.", ["Provide an indicator table: baseline, target, achieved, % of target.", "Disaggregate people reached by sex."]),
  ("Water Supply Component", "Report on water point rehabilitation, water quality and water user committees.", ["How many water points were rehabilitated and how many people gained access?", "Did water quality meet the chlorine standard?"]),
  ("Sanitation and Hygiene Component", "Report on latrines, school WASH blocks and hygiene promotion.", ["How many latrines and school blocks were completed?", "What changed in handwashing practice?"]),
  ("Financial Summary", "Provide budget versus actual expenditure by budget line.", ["Include a table: budget line, budget, spent, balance."]),
  ("Challenges and Lessons Learned", "Describe key challenges, how they were addressed and lessons for future programming.", []),
  ("Sustainability and Exit", "Explain how results will be sustained after the project ends.", ["How are water user committees prepared to maintain the schemes?"]),
  ("Conclusion and Recommendations", "Conclude and recommend next steps.", [])],
 "Submit within 30 days of project end.")
