#!/usr/bin/env python3
"""Build the 20-slide DonorDesk BOARD-APPROVAL deck.

Audience: board members and non-technical approvers. This deck sells the vision,
the outcomes and the business case. Mechanisms, module lists and contracts live
in the technical deck (scripts/build-donordesk-deck.py) — not here.

Colour scheme = the commercial site's design tokens (single source of truth):
  apps/web/tailwind.config.ts   -> brand / accent / success / warning / ai scales
  apps/web/src/app/globals.css  -> --background, --surface, --text, ambient glows
Copy (feature names, tags, how-it-works steps, pricing, security lines) mirrors
  apps/web/src/app/page.tsx  (published at https://donordesk.online)

Usage:
    pip3 install --target /tmp/pptxlib python-pptx
    PYTHONPATH=/tmp/pptxlib python3 scripts/build-donordesk-board-deck.py \
        --out DonorDesk-Board-Briefing-v2.0.pptx

Facts as of 2026-09-28 from memorybank/: base/DonorDesk — Initial Concept
Document.md, base/DonorDesk — One-Page Concept Note for Approval.md,
imp/DonorDesk — Phased Implementation Plan.md, imp/AI-REPORTER-2-RESULTS.md,
imp/Phase21-agent-memory.md, Features/, pending.md.
"""

from __future__ import annotations

import argparse
import math

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# --------------------------------------------------------------------------- #
# Theme — DonorDesk brand tokens (tailwind key -> hex)
# --------------------------------------------------------------------------- #

FONT = "Calibri"
MONO = "Consolas"

INK = "0F172A"        # globals.css --text
BODY = "334155"       # body copy
MUTED = "64748B"      # captions
FAINT = "94A3B8"      # globals.css --text-muted
BRAND = "0C8DE6"      # brand-500
BRAND_DK = "0070C4"   # brand-600 (primary buttons)
BRAND_DP = "015AA1"   # brand-700
BRAND_900 = "0B3F6F"  # brand-900
CYAN = "06B6D4"       # accent-500
CYAN_LT = "22D3EE"    # accent-400
BRAND_LT = "7CC7FB"   # brand-300 (accent text on dark surfaces)
DARK = "020617"       # .dark --background
DARK_CHIP = "12233A"  # raised surface on dark
DARK_LINE = "24405F"
BG = "EEF2F8"         # globals.css --background (light theme)
SURFACE = "FFFFFF"    # --surface-strong
TINT = "F0F7FF"       # brand-50
TINT2 = "E0EFFE"      # brand-100
BORDER = "DCE3ED"     # hairline
SUCCESS = "16A34A"    # success-600
AMBER = "D97706"      # warning-600
RED = "DC2626"        # danger-600
AI = "7C3AED"         # ai-600

SW = 13.333           # slide width (in)
SH = 7.5              # slide height (in)
ML = 0.62             # left margin
MR = 0.62             # right margin
CW = SW - ML - MR     # content width
TOP = 1.42            # default content top
BOTTOM = 6.82         # content bottom
CH = BOTTOM - TOP     # default content height

FOOTER_LEFT = "DonorDesk · Evidence intelligence for funded programmes"


# --------------------------------------------------------------------------- #
# Primitives
# --------------------------------------------------------------------------- #

def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def rect(slide, x, y, w, h, fill=None, line=None, radius=None, line_w=1.0):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None:
        shp.adjustments[0] = radius
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = RGBColor.from_string(line)
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def oval(slide, x, y, w, h, fill=None, alpha=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        shp.fill.background()
    shp.line.fill.background()
    shp.shadow.inherit = False
    if alpha is not None:
        # python-pptx has no alpha API; inject <a:alpha> into the solidFill.
        from pptx.oxml.ns import qn
        solid = shp.fill._xPr.find(qn("a:solidFill"))
        if solid is not None:
            clr = solid.find(qn("a:srgbClr"))
            if clr is not None:
                clr.append(clr.makeelement(qn("a:alpha"),
                                           {"val": str(int(alpha * 1000))}))
    return shp


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, wrap=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, text, size=12, bold=False, color=BODY, first=False, italic=False,
         space_before=0, space_after=4, indent=0.0, line_spacing=1.0,
         align=PP_ALIGN.LEFT, font=FONT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line_spacing
    if indent:
        p.left_indent = Inches(indent)
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = font
    return p


def rich_para(tf, chunks, size=12, color=BODY, first=False, space_before=0,
              space_after=4, indent=0.0, line_spacing=1.0, align=PP_ALIGN.LEFT,
              font=FONT):
    """chunks: list of (text, bold, color|None, size|None)."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line_spacing
    if indent:
        p.left_indent = Inches(indent)
    for text, bold, ch_color, ch_size in chunks:
        run = p.add_run()
        run.text = text
        run.font.size = Pt(ch_size or size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(ch_color or color)
        run.font.name = font
    return p


def shape_text(shape, lines, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT,
               margins=(0.14, 0.14, 0.07, 0.07)):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(margins[0])
    tf.margin_right = Inches(margins[1])
    tf.margin_top = Inches(margins[2])
    tf.margin_bottom = Inches(margins[3])
    for i, spec in enumerate(lines):
        para(tf, spec.get("text", ""), size=spec.get("size", 11),
             bold=spec.get("bold", False), color=spec.get("color", BODY),
             first=(i == 0), space_after=spec.get("space_after", 2),
             space_before=spec.get("space_before", 0),
             align=spec.get("align", align), italic=spec.get("italic", False),
             line_spacing=spec.get("line_spacing", 1.0),
             font=spec.get("font", FONT))
    return tf


def footer(slide, number, note=None):
    tf = textbox(slide, ML, 6.99, CW * 0.80, 0.28)
    para(tf, note or FOOTER_LEFT, size=8.5, color=FAINT, first=True, space_after=0)
    tf2 = textbox(slide, SW - MR - 1.3, 6.99, 1.3, 0.28)
    para(tf2, f"{number} / 20", size=8.5, color=FAINT, first=True, space_after=0,
         align=PP_ALIGN.RIGHT)


def bg(slide, dark=False, glow=True):
    """Site-like page: --background fill plus the ambient brand/cyan glow."""
    rect(slide, 0, 0, SW, SH, fill=DARK if dark else BG)
    if glow:
        if dark:
            oval(slide, 8.1, -2.2, 7.8, 7.8, fill=BRAND_DK, alpha=32)
            oval(slide, 9.8, 3.0, 6.4, 6.4, fill=CYAN, alpha=16)
            oval(slide, -2.8, 3.6, 6.6, 6.6, fill=BRAND, alpha=14)
        else:
            oval(slide, 8.7, -2.5, 7.4, 7.4, fill=BRAND, alpha=14)
            oval(slide, -2.9, 4.1, 6.8, 6.8, fill=CYAN, alpha=10)
    return slide


def hero(slide, number, kicker, title, lede=None, note=None, title_size=24):
    """Standard page header. Returns the y where slide content may start."""
    y = 0.34
    tf = textbox(slide, ML, y, CW, 0.24)
    para(tf, kicker.upper(), size=9.5, bold=True, color=BRAND, first=True,
         space_after=0)
    tf = textbox(slide, ML, y + 0.26, CW, 0.46)
    para(tf, title, size=title_size, bold=True, color=INK, first=True,
         space_after=0)
    y += 0.76
    if lede:
        tf = textbox(slide, ML, y, CW * 0.88, 0.5)
        para(tf, lede, size=11.5, color=BODY, first=True, space_after=0,
             line_spacing=1.10)
        y += 0.20 * max(1, math.ceil(len(lede) / 112)) + 0.06
    rect(slide, ML, y + 0.04, 1.35, 0.055, fill=BRAND)
    rect(slide, ML + 1.35, y + 0.04, 0.30, 0.055, fill=CYAN)
    footer(slide, number, note)
    return y + 0.28


# --------------------------------------------------------------------------- #
# Composite layouts
# --------------------------------------------------------------------------- #

def bullet_list(slide, items, x=ML, y=TOP, w=CW, size=12.5, sub_size=11,
                bullet_color=BRAND, line_spacing=1.08, space_after=6):
    """items: (level, text) or (level, text, color); level 0 = bullet, 1 = sub."""
    tf = textbox(slide, x, y, w, CH)
    first = True
    for item in items:
        level, text = item[0], item[1]
        color = item[2] if len(item) > 2 else None
        if level == 0:
            rich_para(tf, [("▪  ", True, bullet_color, size),
                           (text, False, color or BODY, size)],
                      first=first, space_after=space_after, line_spacing=line_spacing)
        else:
            rich_para(tf, [("–  ", False, MUTED, sub_size),
                           (text, False, MUTED, sub_size)],
                      first=first, space_after=space_after - 1, indent=0.30,
                      line_spacing=line_spacing)
        first = False
    return tf


def itemlist(slide, x, y, w, items, glyph="✓", color=SUCCESS, size=11.5,
             space_after=7, line_spacing=1.06):
    """items: str, or (text, color), or (glyph, text, color)."""
    tf = textbox(slide, x, y, w, CH)
    first = True
    for item in items:
        if isinstance(item, tuple):
            if len(item) == 3:
                g, text, col = item
            else:
                g, text, col = glyph, item[0], item[1]
        else:
            g, text, col = glyph, item, color
        rich_para(tf, [(f"{g}  ", True, col, size), (text, False, BODY, size)],
                  first=first, space_after=space_after, line_spacing=line_spacing)
        first = False
    return tf


def cards(slide, items, x=ML, y=TOP, w=CW, h=CH, cols=3, gap=0.20,
          title_size=12, body_size=10.5, accent=BRAND, rows=None, fill=SURFACE):
    """items: dict(title=, body=[lines] or str, accent=, badge=, fill=)."""
    n = len(items)
    rows = rows or (n + cols - 1) // cols
    cw = (w - gap * (cols - 1)) / cols
    chh = (h - gap * (rows - 1)) / rows
    for i, item in enumerate(items):
        r, c = divmod(i, cols)
        cx = x + c * (cw + gap)
        cy = y + r * (chh + gap)
        card = rect(slide, cx, cy, cw, chh, fill=item.get("fill", fill),
                    line=BORDER, radius=0.06)
        rect(slide, cx, cy, 0.055, chh, fill=item.get("accent", accent),
             radius=0.30)
        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.18)
        tf.margin_right = Inches(0.14)
        tf.margin_top = Inches(0.13)
        tf.margin_bottom = Inches(0.10)
        title = item["title"]
        if item.get("badge"):
            title = f"{item['badge']}  {title}"
        para(tf, title, size=title_size, bold=True, color=INK, first=True,
             space_after=4, line_spacing=0.98)
        body = item.get("body", [])
        if isinstance(body, str):
            body = [body]
        for line in body:
            text, kind = line if isinstance(line, tuple) else (line, "body")
            if kind == "sub":
                para(tf, text, size=body_size - 0.5, color=MUTED, space_after=2,
                     indent=0.16, line_spacing=1.02)
            elif kind == "mono":
                para(tf, text, size=body_size - 0.5, color=BRAND_DP, space_after=3,
                     font=MONO, line_spacing=1.15)
            elif kind == "head":
                para(tf, text, size=body_size, bold=True, color=INK, space_after=2,
                     space_before=4)
            else:
                para(tf, text, size=body_size,
                     color=item.get("body_color", BODY), space_after=3,
                     line_spacing=1.06)
    return None


def table(slide, headers, rows, x=ML, y=TOP, w=CW, col_widths=None,
          size=10, head_size=10.5, row_h=0.34, head_h=0.36, head_fill=BRAND_DP,
          zebra=TINT):
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x),
                                  Inches(y), Inches(w), Inches(head_h))
    tbl = shape.table
    tbl.first_row = False
    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Inches(w * cw / total)
    for c, text in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(head_fill)
        cell.margin_left = cell.margin_right = Inches(0.09)
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        para(cell.text_frame, text, size=head_size, bold=True, color=SURFACE,
             first=True, space_after=0)
    for r, row in enumerate(rows, start=1):
        tbl.rows[r].height = Inches(row_h)
        for c, value in enumerate(row):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(
                SURFACE if r % 2 else zebra)
            cell.margin_left = cell.margin_right = Inches(0.09)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            para(cell.text_frame, value, size=size, bold=(c == 0),
                 color=INK if c == 0 else BODY, first=True, space_after=0,
                 line_spacing=0.98)
    tbl.rows[0].height = Inches(head_h)
    return tbl


def stat_tile(slide, x, y, w, h, value, label, sub=None, color=BRAND_DK,
              fill=SURFACE, value_size=26, label_size=10):
    box = rect(slide, x, y, w, h, fill=fill, line=BORDER, radius=0.07)
    rect(slide, x, y, w, 0.065, fill=color, radius=0.35)
    lines = [
        {"text": value, "size": value_size, "bold": True, "color": color,
         "space_after": 2, "align": PP_ALIGN.CENTER, "line_spacing": 0.95},
        {"text": label, "size": label_size, "color": BODY, "line_spacing": 1.02,
         "align": PP_ALIGN.CENTER, "space_after": 1},
    ]
    if sub:
        lines.append({"text": sub, "size": 8.5, "color": FAINT, "space_after": 0,
                      "align": PP_ALIGN.CENTER, "line_spacing": 1.0})
    shape_text(box, lines, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER,
               margins=(0.10, 0.10, 0.06, 0.06))
    return box


def stat_row(slide, items, x=ML, y=TOP, w=CW, h=1.25, gap=0.18):
    """items: (value, label, color) or (value, label, color, sub)."""
    n = len(items)
    bw = (w - gap * (n - 1)) / n
    for i, item in enumerate(items):
        value, label, color = item[0], item[1], item[2]
        sub = item[3] if len(item) > 3 else None
        stat_tile(slide, x + i * (bw + gap), y, bw, h, value, label, sub, color)


def chips(slide, items, x=ML, y=TOP, w=CW, h=0.36, gap=0.14, size=10,
          color=BRAND_DP, fill=TINT, border=None):
    """items: str or (str, color). Evenly distributed pills."""
    n = len(items)
    bw = (w - gap * (n - 1)) / n
    for i, item in enumerate(items):
        text, col = item if isinstance(item, tuple) else (item, color)
        box = rect(slide, x + i * (bw + gap), y, bw, h, fill=fill,
                   line=border, radius=0.45)
        shape_text(box, [{"text": text, "size": size, "bold": True,
                          "color": col, "align": PP_ALIGN.CENTER,
                          "space_after": 0}],
                   anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER,
                   margins=(0.10, 0.10, 0.02, 0.02))


def quote(slide, text, x=ML, y=TOP, w=CW, h=0.85, fill=BRAND_900, accent=CYAN,
          size=13.5, color=SURFACE):
    box = rect(slide, x, y, w, h, fill=fill, radius=0.06)
    rect(slide, x, y, 0.07, h, fill=accent)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.26)
    tf.margin_right = Inches(0.22)
    para(tf, text, size=size, bold=True, color=color, first=True, space_after=0,
         line_spacing=1.06)
    return box


def panel(slide, x, y, w, h, title, lines, accent=BRAND, fill=SURFACE,
          title_size=12.5, body_size=11):
    """lines: str or (text, kind) with kind in body/sub/mono/head/item/warn."""
    box = rect(slide, x, y, w, h, fill=fill, line=BORDER, radius=0.06)
    rect(slide, x, y, w, 0.055, fill=accent, radius=0.30)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.20)
    tf.margin_right = Inches(0.16)
    tf.margin_top = Inches(0.18)
    tf.margin_bottom = Inches(0.12)
    para(tf, title, size=title_size, bold=True, color=INK, first=True,
         space_after=6, line_spacing=0.98)
    for line in lines:
        text, kind = line if isinstance(line, tuple) else (line, "body")
        if kind == "sub":
            para(tf, text, size=body_size - 0.5, color=MUTED, space_after=2,
                 indent=0.18, line_spacing=1.04)
        elif kind == "mono":
            para(tf, text, size=body_size - 0.5, color=BRAND_DP, space_after=2,
                 font=MONO, line_spacing=1.22)
        elif kind == "head":
            para(tf, text, size=body_size, bold=True, color=INK, space_after=2,
                 space_before=5)
        elif kind == "item":
            rich_para(tf, [("✓  ", True, SUCCESS, body_size),
                           (text, False, BODY, body_size)], space_after=4,
                      line_spacing=1.06)
        elif kind == "warn":
            rich_para(tf, [("▲  ", True, AMBER, body_size),
                           (text, False, BODY, body_size)], space_after=4,
                      line_spacing=1.06)
        else:
            para(tf, text, size=body_size, color=BODY, space_after=4,
                 line_spacing=1.06)
    return box


def progress(slide, x, y, w, h, pct, color=BRAND, track=TINT2, label=None,
             label_size=9.5):
    rect(slide, x, y, w, h, fill=track, radius=0.45)
    filled = max(0.12, w * pct / 100.0)
    rect(slide, x, y, filled, h, fill=color, radius=0.45)
    if pct >= 24:
        tf = textbox(slide, x, y, filled, h, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, label or f"{pct}% ready", size=label_size, bold=True,
             color=SURFACE, first=True, space_after=0, align=PP_ALIGN.CENTER)
    else:
        tf = textbox(slide, x + filled + 0.08, y, w * 0.6, h,
                     anchor=MSO_ANCHOR.MIDDLE)
        para(tf, label or f"{pct}% ready", size=label_size, bold=True,
             color=color, first=True, space_after=0)


def strip(slide, text, y=6.26, x=ML, w=CW, h=0.44, fill=BRAND_900,
          color=SURFACE, accent=CYAN, size=10.5, highlight=None, border=None):
    box = rect(slide, x, y, w, h, fill=fill, line=border, radius=0.08)
    rect(slide, x, y, 0.07, h, fill=accent)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.24)
    tf.margin_right = Inches(0.20)
    if highlight:
        rich_para(tf, [(text + "  ", True, color, size),
                       (highlight, True, CYAN_LT, size)], first=True,
                  space_after=0, line_spacing=1.05)
    else:
        para(tf, text, size=size, color=color, first=True, space_after=0,
             line_spacing=1.05)
    return box


def timeline(slide, phases, x=ML, y=TOP, w=CW, h=CH, gap=0.20):
    """phases: dict(kicker=, title=, bullets=[], exit=, accent=, fill=)."""
    n = len(phases)
    cw = (w - gap * (n - 1)) / n
    rect(slide, x, y + 0.30, w, 0.032, fill=TINT2)
    for i, ph in enumerate(phases):
        cx = x + i * (cw + gap)
        accent = ph.get("accent", BRAND)
        dark = ph.get("dark", False)
        head = rect(slide, cx, y, cw, 0.80, fill=accent, radius=0.07)
        shape_text(head, [
            {"text": ph["kicker"], "size": 9, "bold": True,
             "color": INK if accent == CYAN_LT else SURFACE, "space_after": 3},
            {"text": ph["title"], "size": 12.5, "bold": True,
             "color": INK if accent == CYAN_LT else SURFACE, "space_after": 0,
             "line_spacing": 0.98},
        ], anchor=MSO_ANCHOR.MIDDLE)
        body_y = y + 0.94
        box = rect(slide, cx, body_y, cw, h - (body_y - y) - 0.70,
                   fill=DARK_CHIP if dark else ph.get("fill", SURFACE),
                   line=DARK_LINE if dark else BORDER, radius=0.06)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.17)
        tf.margin_right = Inches(0.14)
        tf.margin_top = Inches(0.14)
        tf.margin_bottom = Inches(0.10)
        first = True
        for b in ph["bullets"]:
            rich_para(tf, [("▪  ", True, accent if not dark else CYAN_LT, 10.5),
                           (b, False, SURFACE if dark else BODY, 10.5)],
                      first=first, space_after=5, line_spacing=1.06)
            first = False
        exit_box = rect(slide, cx, y + h - 0.62, cw, 0.62,
                        fill=ph.get("exit_fill", TINT), line=BORDER, radius=0.07)
        shape_text(exit_box, [
            {"text": "EXIT TEST", "size": 8, "bold": True, "color": accent,
             "space_after": 2},
            {"text": ph["exit"], "size": 9.5, "color": INK, "space_after": 0,
             "line_spacing": 1.03},
        ], anchor=MSO_ANCHOR.MIDDLE)
    return None


def code_card(slide, x, y, w, h, title, lines, accent=BRAND, fill=SURFACE,
              size=10, title_size=11, note=None):
    box = rect(slide, x, y, w, h, fill=fill, line=BORDER, radius=0.06)
    rect(slide, x, y, w, 0.055, fill=accent, radius=0.30)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.20)
    tf.margin_right = Inches(0.16)
    tf.margin_top = Inches(0.18)
    tf.margin_bottom = Inches(0.12)
    para(tf, title, size=title_size, bold=True, color=INK, first=True,
         space_after=6, line_spacing=0.98)
    for line in lines:
        if isinstance(line, tuple):
            text, kind = line
        else:
            text, kind = line, "mono"
        if kind == "note":
            para(tf, text, size=size - 0.5, color=MUTED, space_after=3,
                 line_spacing=1.06)
        else:
            para(tf, text, size=size, color=BRAND_DP, font=MONO, space_after=2,
                 line_spacing=1.22)
    if note:
        para(tf, note, size=8.5, color=FAINT, space_after=0, space_before=5)
    return box


# --------------------------------------------------------------------------- #
# Slide 1 — Cover
# --------------------------------------------------------------------------- #

def s01_title(prs):
    slide = bg(blank(prs), dark=True)

    tf = textbox(slide, ML, 1.22, 8.4, 0.3)
    para(tf, "BOARD BRIEFING  ·  VERSION 2.0  ·  SEPTEMBER 2026", size=10.5,
         bold=True, color=CYAN_LT, first=True, space_after=0)

    tf = textbox(slide, ML, 1.58, 8.4, 0.9)
    para(tf, "DonorDesk", size=46, bold=True, color=SURFACE, first=True,
         space_after=0)

    rect(slide, ML, 2.42, 1.5, 0.06, fill=BRAND_LT)
    rect(slide, ML + 1.5, 2.42, 0.34, 0.06, fill=CYAN)

    tf = textbox(slide, ML, 2.64, 8.0, 0.5)
    para(tf, "Evidence intelligence for funded programmes.", size=21, bold=True,
         color=BRAND_LT, first=True, space_after=0)

    tf = textbox(slide, ML, 3.24, 7.9, 0.5)
    para(tf, "From scattered field evidence to donor-ready reports.", size=15,
         color="E2E8F0", first=True, space_after=0)

    tf = textbox(slide, ML, 3.78, 7.7, 1.0)
    para(tf, "DonorDesk helps humanitarian, development and other grant-funded "
             "programmes turn activity records, indicator results and supporting "
             "evidence into professional, audit-ready reports — with "
             "source-linked AI drafting, automatic compliance checks and human "
             "approval.", size=12, color=FAINT, first=True, space_after=0,
         line_spacing=1.18)

    chips(slide, ["Human-reviewed AI", "Source-linked drafting",
                  "Live readiness", "Audit-ready exports"],
          x=ML, y=5.05, w=7.7, h=0.40, gap=0.12, size=10, color=CYAN_LT,
          fill=DARK_CHIP, border=DARK_LINE)

    # site-style readiness panel (illustrative content, published online)
    px, py, pw, ph = 8.66, 1.58, 4.05, 3.87
    box = rect(slide, px, py, pw, ph, fill=DARK_CHIP, line=DARK_LINE, radius=0.05)
    rect(slide, px, py, pw, 0.06, fill=BRAND)
    shape_text(box, [
        {"text": "Emergency Response — Health Programme", "size": 10.5,
         "bold": True, "color": SURFACE, "space_after": 2},
        {"text": "Report readiness · Q3 2026", "size": 9, "color": FAINT,
         "space_after": 8},
        {"text": "92% ready", "size": 22, "bold": True, "color": CYAN_LT,
         "space_after": 6},
    ], anchor=MSO_ANCHOR.TOP, margins=(0.20, 0.18, 0.18, 0.10))
    progress(slide, px + 0.20, py + 1.32, pw - 0.40, 0.20, 92, color=BRAND,
             track=DARK_LINE, label="")
    tf = textbox(slide, px + 0.20, py + 1.72, pw - 0.40, 1.9)
    for i, line in enumerate([
        ("✓", "Evidence verified", "128 / 132", SUCCESS),
        ("✓", "Indicators on track", "24 / 25", SUCCESS),
        ("✓", "Activities logged", "61 / 64", SUCCESS),
        ("▲", "Indicators awaiting verification", "2", AMBER),
    ]):
        rich_para(tf, [(f"{line[0]}  ", True, line[3], 10.5),
                       (line[1], False, "CBD5E1", 10.5),
                       (f"   {line[2]}", True, SURFACE, 10.5)],
                  first=(i == 0), space_after=8, line_spacing=1.0)
    tf = textbox(slide, px + 0.20, py + ph - 0.42, pw - 0.40, 0.3)
    para(tf, "Illustrative view, as published on donordesk.online", size=8,
         color=FAINT, first=True, space_after=0)

    tf = textbox(slide, ML, 6.28, 6.0, 0.3)
    para(tf, "donordesk.online", size=12, bold=True, color=CYAN_LT, first=True,
         space_after=0)
    tf = textbox(slide, SW - MR - 5.0, 6.28, 5.0, 0.3)
    para(tf, "Prepared for board review · 2026-09-28", size=10, color=FAINT,
         first=True, space_after=0, align=PP_ALIGN.RIGHT)


# --------------------------------------------------------------------------- #
# Slide 2 — Executive summary
# --------------------------------------------------------------------------- #

def s02_exec_summary(prs):
    slide = bg(blank(prs))
    y = hero(
        slide, 2, "EXECUTIVE SUMMARY", "Where we are, in one page",
        "DonorDesk turns programme delivery records and field evidence into "
        "donor-ready reports — with the figures owned by the system, a source "
        "behind every claim, and a human on the approve button.")

    stat_row(slide, [
        ("21", "product features delivered", BRAND_DK, "Features 01–21"),
        ("Live", "in production today", SUCCESS, "donordesk.online · HTTPS"),
        ("25 / 25", "report-quality acceptance corpus", BRAND_DP,
         "every hard gate green"),
        ("$0–$149", "published monthly pricing", AI, "plus Enterprise"),
    ], y=y, h=1.34)

    panel(slide, ML, y + 1.56, CW / 2 - 0.11, 2.34, "What the product does", [
        ("Build the reporting structure once — logframe, indicators, donor "
         "template, compliance requirements.", "item"),
        ("Capture activities, results and evidence as delivery happens.", "item"),
        ("Draft, check, review and export every reporting period from those "
         "records.", "item"),
    ], accent=BRAND, body_size=11.5)

    panel(slide, ML + CW / 2 + 0.11, y + 1.56, CW / 2 - 0.11, 2.34,
          "What the platform guarantees", [
        ("Every number comes from verified programme records — the model writes "
         "prose, the system owns the figures.", "item"),
        ("Every AI output is reviewable: model, prompt version and sources are "
         "stored with it.", "item"),
        ("Tenant isolation is enforced in the database, and every change is "
         "audited.", "item"),
    ], accent=CYAN, body_size=11.5)

    strip(slide, "Board decisions requested today:",
          highlight="approve launch pricing and the nonprofit policy, and back a "
                    "3–5 NGO pilot programme.",
          y=6.10, h=0.50)


# --------------------------------------------------------------------------- #
# Slide 3 — The problem
# --------------------------------------------------------------------------- #

def s03_problem(prs):
    slide = bg(blank(prs))
    y = hero(
        slide, 3, "THE PROBLEM",
        "Reporting is the tax every funded programme pays",
        "The work gets done in the field. Proving it happens later, by hand — "
        "reassembling a period from WhatsApp threads, inboxes, spreadsheets and "
        "field folders while a donor deadline moves closer.")

    cards(slide, [
        {
            "badge": "01", "title": "Evidence is scattered",
            "accent": AMBER,
            "body": [
                "Attendance sheets, distribution lists, photos, approvals and "
                "field notes live in WhatsApp, email, personal drives and shared "
                "folders.",
                ("Nothing is linked to the activity, indicator or donor "
                 "requirement it is meant to prove.", "sub"),
                "The same attendance sheet gets emailed to three different "
                "people.",
            ],
        },
        {
            "badge": "02", "title": "Quality is invisible until deadline week",
            "accent": RED,
            "body": [
                "Gaps surface during final assembly: an expired approval, an "
                "unexplained indicator movement, an annex nobody owns.",
                ("By then the only fixes available are late ones.", "sub"),
                "Readiness cannot be seen until somebody assembles the document.",
            ],
        },
        {
            "badge": "03", "title": "Knowledge leaves when staff rotate",
            "accent": BRAND,
            "body": [
                "Emergency and multi-donor programmes rotate teams every few "
                "months, so every cycle rebuilds the same story from scratch.",
                ("Institutional memory ends up as a folder nobody can navigate.",
                 "sub"),
                "New staff start by asking where last period's files are.",
            ],
        },
    ], y=y, h=2.86, cols=3, gap=0.22, title_size=13, body_size=11.5)

    quote(slide, "The work was done. Proving it happened took longer than doing "
                 "it.", y=y + 3.04, h=0.72, size=15)

    strip(slide, "What it costs:",
          highlight="weeks of assembly every cycle · late annexes · audit "
                    "exposure · donor confidence.",
          y=y + 3.94, h=0.46)


# --------------------------------------------------------------------------- #
# Slide 4 — Who it is for
# --------------------------------------------------------------------------- #

def s04_who(prs):
    slide = bg(blank(prs))
    y = hero(
        slide, 4, "WHO IT IS FOR",
        "Built first for humanitarian and development teams",
        "Designed for any grant-funded programme that must prove results — "
        "multiple funders, multiple locations, one evidence trail.")

    personas = [
        ("Programme Manager", "Sees delivery and gaps in one place instead of "
                              "chasing updates before a deadline.", BRAND),
        ("M&E Officer", "Indicators, results and evidence reconcile instead of "
                        "living in a spreadsheet beside the narrative.", CYAN),
        ("Grants & Reporting Lead", "Starts from a structure that already knows "
                                    "the funder's template and requirements.", BRAND_DK),
        ("Compliance & Audit", "Every claim traces to a source, a version and an "
                               "approval in the audit trail.", BRAND_DP),
        ("Country Director", "Predictable submission quality across donors, "
                             "teams and locations.", AI),
    ]
    lw = 7.66
    row_h, gap = 0.78, 0.10
    for i, (name, desc, accent) in enumerate(personas):
        cy = y + i * (row_h + gap)
        box = rect(slide, ML, cy, lw, row_h, fill=SURFACE, line=BORDER,
                   radius=0.06)
        rect(slide, ML, cy, 0.055, row_h, fill=accent, radius=0.30)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.18)
        tf.margin_right = Inches(0.14)
        para(tf, name, size=11.5, bold=True, color=INK, first=True, space_after=2,
             line_spacing=0.98)
        para(tf, desc, size=10.5, color=MUTED, space_after=0, line_spacing=1.04)

    sx = ML + lw + 0.22
    sw = CW - lw - 0.22
    panel(slide, sx, y, sw, row_h * 2 + gap + 0.36, "Sectors we serve today", [
        "◎  Humanitarian response",
        "✦  International development",
        "◉  Public health",
        "▤  Research & education",
        "◇  Government grants",
        "≈  Climate & social impact",
    ], accent=CYAN, body_size=11.5)

    panel(slide, sx, y + row_h * 2 + gap + 0.58, sw,
          row_h * 3 + gap * 2 - 0.58, "Why it lands with donors too", [
        ("Submissions arrive complete, consistent and on time.", "item"),
        ("Claims are substantiated, so review questions drop.", "item"),
        ("The audit trail exists before it is ever requested.", "item"),
    ], accent=SUCCESS, body_size=11.5)


# --------------------------------------------------------------------------- #
# Slide 5 — The vision
# --------------------------------------------------------------------------- #

def s05_vision(prs):
    slide = bg(blank(prs))
    y = hero(
        slide, 5, "THE VISION",
        "The operating layer for donor-funded reporting",
        "DonorDesk should become the system an organisation opens when it must "
        "answer — for any programme, any funder, any period:")

    questions = [
        "What activities were completed?",
        "Where did they take place?",
        "Who benefited?",
        "Which indicators changed?",
        "What evidence proves it?",
        "Which donor requirement does it satisfy?",
        "What is still missing before submission?",
        "Which compliance risk needs attention?",
    ]
    cards(slide, [
        {"badge": "›", "title": q, "accent": BRAND if i < 4 else CYAN,
         "body": []}
        for i, q in enumerate(questions)
    ], y=y, h=2.66, cols=4, rows=2, gap=0.18, title_size=12.5, body_size=10)

    cards(slide, [
        {
            "badge": "01", "title": "Own the evidence",
            "accent": BRAND, "body": [
                "One structured place where delivery, results, documents and "
                "approvals are linked to the programme structure.",
            ],
        },
        {
            "badge": "02", "title": "Own the reporting cycle",
            "accent": CYAN, "body": [
                "Every period repeats the same auditable path: capture, draft, "
                "check, review, approve, export.",
            ],
        },
        {
            "badge": "03", "title": "Own the standard",
            "accent": AI, "body": [
                "Sector reporting packs and donor requirements become the "
                "reason teams choose DonorDesk first.",
            ],
        },
    ], y=y + 2.84, h=1.62, cols=3, gap=0.18, title_size=12, body_size=10.5)


# --------------------------------------------------------------------------- #
# Slide 6 — How it works
# --------------------------------------------------------------------------- #

def s06_how(prs):
    slide = bg(blank(prs))
    y = hero(
        slide, 6, "HOW IT WORKS",
        "One clear path from set-up to periodic reporting",
        "Build the reporting structure once, keep it current as delivery unfolds, "
        "and turn verified programme records into repeatable reports.")

    cards(slide, [
        {"badge": "01", "title": "Set up your project",
         "accent": BRAND,
         "body": ["Define the programme, reporting profile, dates, funder and "
                  "team."]},
        {"badge": "02", "title": "Build your logframe",
         "accent": BRAND,
         "body": ["Create it in DonorDesk, or import objectives, outcomes and "
                  "outputs."]},
        {"badge": "03", "title": "Add your indicators",
         "accent": CYAN,
         "body": ["Enter or import baselines, targets, results and data "
                  "sources."]},
        {"badge": "04", "title": "Define compliance",
         "accent": CYAN,
         "body": ["Turn funder requirements into a live, trackable readiness "
                  "checklist."]},
        {"badge": "05", "title": "Capture delivery & evidence",
         "accent": BRAND_DK,
         "body": ["Log activities, results, documents, photos and supporting "
                  "data as work happens."]},
        {"badge": "06", "title": "Generate ready-to-review reports",
         "accent": BRAND_DK,
         "body": ["Draft periodically, verify sources, approve sections and "
                  "export funder-ready files."]},
    ], y=y, h=3.56, cols=3, rows=2, gap=0.20, title_size=12.5, body_size=10.5)

    strip(slide, "The chain that keeps it honest:",
          highlight="logframe → indicators → activities → evidence → compliance "
                    "→ draft → review → export",
          y=5.66, h=0.52)


# --------------------------------------------------------------------------- #
# Slide 7 — Before and after
# --------------------------------------------------------------------------- #

def s07_before_after(prs):
    slide = bg(blank(prs))
    y = hero(
        slide, 7, "THE DIFFERENCE",
        "The same monthly report, two very different weeks",
        "Illustrative walkthrough: a funded health programme reporting on a "
        "single month.")

    gap = 0.38
    pw = (CW - gap) / 2
    h = 3.52

    panel(slide, ML, y, pw, h, "Today — assembled by hand", [
        "Activity updates arrive late; nobody knows what is still missing.",
        "M&E cleans indicator data in a spreadsheet beside the narrative.",
        "The story is rewritten from emails and last period's report.",
        "Annexes are chased from people; the donor template is rebuilt in Word.",
        "Reporting weeks cannot be planned — nobody can see what is "
        "outstanding.",
        ("The report is finished in the final week, under pressure.", "sub"),
    ], accent=AMBER, body_size=11)

    tf = textbox(slide, ML + pw, y + h / 2 - 0.3, gap, 0.6,
                 anchor=MSO_ANCHOR.MIDDLE)
    para(tf, "→", size=26, bold=True, color=BRAND, first=True, space_after=0,
         align=PP_ALIGN.CENTER)

    panel(slide, ML + pw + gap, y, pw, h, "With DonorDesk — assembled by the "
          "system", [
        ("Activities, results and evidence are captured as delivery happens.",
         "item"),
        ("Evidence is already mapped to the activity, indicator and donor "
         "requirement.", "item"),
        ("A first draft appears with a source attached to every statement.",
         "item"),
        ("Readiness names exactly what is missing — days before the deadline.",
         "item"),
        ("PDF, DOCX, XLSX and ZIP produce as one approved pack.", "item"),
    ], accent=BRAND, body_size=11)

    quote(slide, "Less report assembly. More reporting confidence.", y=5.62,
          h=0.62, size=15)


# --------------------------------------------------------------------------- #
# Slide 8 — Feature 1: donor template extraction
# --------------------------------------------------------------------------- #

def s08_templates(prs):
    slide = bg(blank(prs))
    y = hero(
        slide, 8, "SIGNATURE FEATURE 1", "Upload the funder's template once",
        "DonorDesk structures donor and funder templates into editable, "
        "reviewable report sections with evidence requirements mapped to your "
        "programme.")

    lw = 5.52
    panel(slide, ML, y, lw, 3.86, "What changes for the team", [
        ("Stop rebuilding the donor's Word template every reporting cycle.",
         "item"),
        ("Sections, hierarchy, mandatory questions, required tables and page "
         "limits become part of the plan.", "item"),
        ("Every section carries the evidence the donor expects against it.",
         "item"),
        ("Reviewed once, reused for every period of the grant.", "item"),
    ], accent=BRAND, body_size=11)

    rx = ML + lw + 0.22
    rw = CW - lw - 0.22
    code_card(slide, rx, y, rw, 3.86,
              "Extracted structure — illustrative example", [
                  "USAID Performance Report (FY2026)",
                  "├ 1. Programme Overview",
                  "├ 2. Results & Indicators      [3 required tables]",
                  "├ 3. Beneficiary Reach         [annex: lists]",
                  "├ 4. Challenges & Lessons      [mandatory question]",
                  "└ 5. Financial Summary         [page limit: 2]",
                  ("Sections, levels, required tables and mandatory questions are "
                   "extracted from the uploaded document, then confirmed by a "
                   "person before use.", "note"),
              ], accent=BRAND_DK, size=10.5)

    strip(slide, "Impact:",
          highlight="days of formatting disappear from every reporting cycle.",
          y=y + 4.02, h=0.48)


# --------------------------------------------------------------------------- #
# Slide 9 — Feature 2: source-linked AI drafting
# --------------------------------------------------------------------------- #

def s09_drafting(prs):
    slide = bg(blank(prs))
    y = hero(
        slide, 9, "SIGNATURE FEATURE 2", "AI as assistant, not author",
        "Drafts are built from saved activities, indicator results and evidence "
        "text — statement-level sources stay visible, and every section remains "
        "editable and human-approved.")

    lw = 5.52
    panel(slide, ML, y, lw, 3.86, "How we keep it honest", [
        ("The model writes prose. Every figure comes from your verified records.",
         "item"),
        ("Each statement keeps the sources behind it in view.", "item"),
        ("Sections stay editable — nothing is released without approval.",
         "item"),
        ("House style is learned per organisation, never across tenants.",
         "item"),
    ], accent=AI, body_size=11)

    rx = ML + lw + 0.22
    rw = CW - lw - 0.22
    box = rect(slide, rx, y, rw, 3.86, fill=SURFACE, line=BORDER, radius=0.06)
    rect(slide, rx, y, rw, 0.055, fill=AI, radius=0.30)
    tf = textbox(slide, rx + 0.20, y + 0.18, rw - 0.38, 0.3)
    para(tf, "What a drafted statement looks like", size=11, bold=True,
         color=INK, first=True, space_after=0)

    sq = rect(slide, rx + 0.20, y + 0.58, rw - 0.38, 1.06, fill=TINT,
              line=BORDER, radius=0.06)
    shape_text(sq, [{
        "text": "“Community health workers delivered 1,240 consultations in Q3 "
                "(Indic 4.1, [evidence-2041]), an increase of 18% over target.”",
        "size": 11.5, "color": INK, "space_after": 0, "line_spacing": 1.18,
    }], anchor=MSO_ANCHOR.MIDDLE, margins=(0.16, 0.14, 0.10, 0.10))

    tf = textbox(slide, rx + 0.20, y + 1.76, rw - 0.38, 0.3)
    para(tf, "Sources attached to that statement", size=9.5, bold=True,
         color=BRAND_DP, first=True, space_after=0)

    tf = textbox(slide, rx + 0.20, y + 2.06, rw - 0.38, 1.4)
    for i, line in enumerate([
        ("✓", "Indic 4.1 — verified result, Q3 2026", SUCCESS),
        ("✓", "evidence-2041 — attendance register, checksum verified", SUCCESS),
        ("✓", "Target 1,050 — approved logframe", SUCCESS),
        ("◦", "18% over target — derived figure, the only one allowed", MUTED),
    ]):
        rich_para(tf, [(f"{line[0]}  ", True, line[2], 10),
                       (line[1], False, BODY, 10)],
                  first=(i == 0), space_after=6, line_spacing=1.04)

    tf = textbox(slide, rx + 0.20, y + 3.48, rw - 0.38, 0.3)
    para(tf, "Illustrative example, consistent with published product copy.",
         size=8, color=FAINT, first=True, space_after=0)

    strip(slide, "The rule:",
          highlight="the model never invents a number — it explains numbers the "
                    "platform owns.",
          y=y + 4.02, h=0.48)


# --------------------------------------------------------------------------- #
# Slide 10 — Feature 3: live compliance checklist
# --------------------------------------------------------------------------- #

def s10_readiness(prs):
    slide = bg(blank(prs))
    y = hero(
        slide, 10, "SIGNATURE FEATURE 3",
        "Readiness you can see before the deadline",
        "A live readiness score surfaces missing evidence, unverified indicators "
        "and late activity updates while there is still time to fix them.")

    lw = 5.52
    panel(slide, ML, y, lw, 3.86, "What the readiness score is made of", [
        ("Required sections completed", "item"),
        ("Indicators updated with results for the period", "item"),
        ("Evidence attached to activities and indicators", "item"),
        ("Compliance documents verified — approvals, agreements, registrations",
         "item"),
        ("Internal review completed", "item"),
        ("Export-ready, approved deliverables", "item"),
    ], accent=CYAN, body_size=11)

    rx = ML + lw + 0.22
    rw = CW - lw - 0.22
    box = rect(slide, rx, y, rw, 3.86, fill=SURFACE, line=BORDER, radius=0.06)
    rect(slide, rx, y, rw, 0.055, fill=CYAN, radius=0.30)
    tf = textbox(slide, rx + 0.20, y + 0.18, rw - 0.38, 0.6)
    para(tf, "Monthly report — Q3 2026", size=12, bold=True, color=INK,
         first=True, space_after=2)
    para(tf, "Donor deadline: 14 October · 9 days remaining", size=9.5,
         color=MUTED, space_after=0)

    progress(slide, rx + 0.20, y + 0.96, rw - 0.40, 0.34, 72, color=BRAND,
             track=TINT2, label="72% ready")

    tf = textbox(slide, rx + 0.20, y + 1.50, rw - 0.38, 0.3)
    para(tf, "Still missing before submission", size=9.5, bold=True,
         color=AMBER, first=True, space_after=0)

    tf = textbox(slide, rx + 0.20, y + 1.82, rw - 0.38, 1.5)
    for i, line in enumerate([
        "3 attendance sheets — July distribution",
        "1 procurement approval — expired",
        "2 indicator explanations — Indic 4.1, Indic 3.2",
        "No unresolved compliance item is left unnamed",
    ]):
        rich_para(tf, [("▲  " if i < 3 else "✓  ", True,
                        AMBER if i < 3 else SUCCESS, 10.5),
                       (line, False, BODY, 10.5)],
                  first=(i == 0), space_after=7, line_spacing=1.04)

    tf = textbox(slide, rx + 0.20, y + 3.42, rw - 0.38, 0.4)
    para(tf, "Illustrative view; readiness is computed from the programme's own "
             "records.", size=8, color=FAINT, first=True, space_after=0)

    strip(slide, "Impact:",
          highlight="gaps are found while they can still be fixed — not in "
                    "deadline week.",
          y=y + 4.02, h=0.48)


# --------------------------------------------------------------------------- #
# Slide 11 — Feature 4: evidence that holds up to audit
# --------------------------------------------------------------------------- #

def s11_evidence(prs):
    slide = bg(blank(prs))
    y = hero(
        slide, 11, "SIGNATURE FEATURE 4",
        "Evidence that holds up to audit",
        "Programme documents, field evidence and source records stay organised "
        "with checksums, verification states and a traceable history.")

    lw = 5.52
    panel(slide, ML, y, lw, 3.86, "What makes evidence audit-ready", [
        ("A checksum is recorded when a file is uploaded", "item"),
        ("Every document has a verification state, not just a folder",
         "item"),
        ("Evidence is linked to the activity, indicator and donor requirement",
         "item"),
        ("Confidential labels and PII warnings for beneficiary data", "item"),
        ("Optional anonymisation of beneficiary lists", "item"),
        ("Version and export history kept for the whole grant", "item"),
    ], accent=SUCCESS, body_size=11)

    rx = ML + lw + 0.22
    rw = CW - lw - 0.22
    code_card(slide, rx, y, rw, 3.86,
              "Evidence pack handed to the donor — illustrative", [
                  "Donor_Package_Q3_2026.zip",
                  "├ 01_Final_Report/           PDF · sealed",
                  "├ 02_Indicator_Table/        XLSX · verified results",
                  "├ 03_Evidence_Checklist/     XLSX · requirement by requirement",
                  "├ 04_Attendance_Sheets/      18 files · checksummed",
                  "└ 05_Photos/                 42 files · tagged to activities",
                  ("Pack contents follow the funder's requirements; approval is "
                   "recorded before the ZIP is sealed.", "note"),
              ], accent=SUCCESS, size=10.5)

    strip(slide, "Impact:",
          highlight="when a donor or auditor asks “show me the source”, the "
                    "answer already exists.",
          y=y + 4.02, h=0.48)


# --------------------------------------------------------------------------- #
# Slide 12 — Feature 5: logframes, indicators, outcome reporting
# --------------------------------------------------------------------------- #

def s12_logframe(prs):
    slide = bg(blank(prs))
    y = hero(
        slide, 12, "SIGNATURE FEATURE 5", "Numbers first, narrative second",
        "Model objectives, outcomes, outputs and indicators with baselines and "
        "targets — then drive the narrative straight from the numbers.")

    lw = 5.52
    panel(slide, ML, y, lw, 3.86, "Why it matters", [
        ("One structure holds the whole programme: goal, outcomes, outputs, "
         "activities.", "item"),
        ("Indicator results update the report itself, not a parallel "
         "spreadsheet.", "item"),
        ("Tables, charts and period-to-period change are built from verified "
         "findings.", "item"),
        ("The logframe annex, the narrative and the donor template share one "
         "source of truth.", "item"),
    ], accent=BRAND_DK, body_size=11)

    rx = ML + lw + 0.22
    rw = CW - lw - 0.22
    table(slide, ["Indicator", "Baseline", "Target", "Q3 result", "Status"], [
        ["4.1 Consultations delivered", "0", "1,050", "1,240", "118% of target"],
        ["3.2 Households reached", "0", "2,000", "1,860", "93% of target"],
        ["2.5 Health workers trained", "12", "40", "40", "Complete"],
    ], x=rx, y=y + 0.04, w=rw, size=9.5, head_size=9.5, row_h=0.38,
        col_widths=[3.0, 1.0, 1.0, 1.1, 1.6])

    panel(slide, rx, y + 1.72, rw, 2.14, "What happens automatically", [
        ("Tables and charts are generated from the same verified findings.",
         "item"),
        ("Movement against the previous period is shown when one exists.",
         "item"),
        ("The narrative explains the movement — it cannot introduce a new "
         "figure.", "item"),
    ], accent=CYAN, body_size=10.5)

    tf = textbox(slide, rx, y + 3.94, rw, 0.3)
    para(tf, "Illustrative figures.", size=8, color=FAINT, first=True,
         space_after=0)

    strip(slide, "Impact:",
          highlight="narrative, tables and annexes can never disagree — they "
                    "share one source.",
          y=y + 4.16, h=0.48)


# --------------------------------------------------------------------------- #
# Slide 13 — Feature 6: donor-ready exports
# --------------------------------------------------------------------------- #

def s13_exports(prs):
    slide = bg(blank(prs))
    y = hero(
        slide, 13, "SIGNATURE FEATURE 6",
        "Deliverables in the formats funders ask for",
        "PDF, DOCX, XLSX and ZIP produce from one structured workspace, rendered "
        "into the donor's own template — with review and approval gates before "
        "release.")

    lw = 5.52
    panel(slide, ML, y, lw, 3.86, "What the team receives", [
        ("PDF, DOCX, XLSX and ZIP from one workspace", "item"),
        ("Output rendered into the donor's own template and headings", "item"),
        ("Internal drafts stay marked as drafts; finals are sealed", "item"),
        ("Every export recorded in history with its version", "item"),
        ("One repeatable pack per funder, every period", "item"),
    ], accent=BRAND_DP, body_size=11)

    rx = ML + lw + 0.22
    rw = CW - lw - 0.22
    cards(slide, [
        {"badge": "PDF", "title": "Final report",
         "accent": BRAND, "body": ["Print-ready, watermarked until approved."]},
        {"badge": "DOCX", "title": "Editable narrative",
         "accent": BRAND, "body": ["For funders who want to leave comments."]},
        {"badge": "XLSX", "title": "Indicator table & checklist",
         "accent": CYAN, "body": ["Results and compliance, as the donor asks."]},
        {"badge": "ZIP", "title": "Full evidence pack",
         "accent": SUCCESS, "body": ["Report, annexes, sheets and photos."]},
    ], x=rx, y=y, w=rw, h=1.58, cols=2, rows=2, gap=0.16, title_size=11,
        body_size=9.5)

    panel(slide, rx, y + 1.74, rw, 2.12, "Approval before release", [
        ("Section reviewer sign-off", "item"),
        ("Final approver sign-off", "item"),
        ("The sealed version records who approved it, and when", "item"),
    ], accent=AI, body_size=10.5)

    strip(slide, "Impact:",
          highlight="one pack per funder, instead of a week of copy-paste.",
          y=y + 4.02, h=0.48)


# --------------------------------------------------------------------------- #
# Slide 14 — Version 2.0 flagship: Agent Memory
# --------------------------------------------------------------------------- #

def s14_memory(prs):
    slide = bg(blank(prs))
    y = hero(
        slide, 14, "VERSION 2.0 FLAGSHIP · AGENT MEMORY",
        "It learns your house style. It never learns your numbers.",
        "Agent Memory turns an organisation's own edits into writing guidance — "
        "tone, structure, recurring phrasing — while figures, dates, targets and "
        "commitments stay out of its reach.")

    lw = 5.52
    half = 1.86
    panel(slide, ML, y, lw, half, "What it learns", [
        ("Tone of voice and register", "item"),
        ("How sections are ordered and opened", "item"),
        ("Recurring phrasing your reviewers keep", "item"),
        ("What reviewers consistently change", "item"),
    ], accent=AI, body_size=10.5)

    panel(slide, ML, y + half + 0.14, lw, 3.86 - half - 0.14,
          "What it never learns", [
        ("Figures, percentages and dates", "item"),
        ("Targets, commitments or donor obligations", "item"),
        ("Beneficiary names or any personal data", "item"),
        ("Anything from another organisation's workspace", "item"),
    ], accent=RED, body_size=10.5)

    rx = ML + lw + 0.22
    rw = CW - lw - 0.22
    cards(slide, [
        {"badge": "LAYER 1", "title": "Extraction filter",
         "accent": BRAND,
         "body": ["Any edit containing a number or a date is discarded before "
                  "memory is even considered."]},
        {"badge": "LAYER 2", "title": "Entity invariant",
         "accent": CYAN,
         "body": ["A writing memory that contains a figure cannot be "
                  "constructed at all — enforced in the domain model."]},
        {"badge": "LAYER 3", "title": "Human approval",
         "accent": SUCCESS,
         "body": ["Nothing becomes active guidance until a person reviews and "
                  "approves it."]},
    ], x=rx, y=y, w=rw, h=2.0, cols=3, rows=1, gap=0.16, title_size=11.5,
        body_size=10)

    panel(slide, rx, y + 2.16, rw, 1.70, "Where it stands today", [
        ("Shipped with Version 2.0; off by default for every tenant.", "item"),
        ("Scoped to one organisation, reviewable and revocable at any time.",
         "item"),
        ("Style only — no cross-tenant learning, ever.", "item"),
    ], accent=AI, body_size=10.5)

    strip(slide, "Why it matters:",
          highlight="from the second reporting period, the first draft already "
                    "sounds like your organisation.",
          y=y + 4.02, h=0.48)


# --------------------------------------------------------------------------- #
# Slide 15 — Responsible AI
# --------------------------------------------------------------------------- #

def s15_responsible_ai(prs):
    slide = bg(blank(prs))
    y = hero(
        slide, 15, "RESPONSIBLE AI", "A human always holds the approve button",
        "AI accelerates drafting. Accountability stays with the organisation — "
        "and every AI output can be reviewed after the fact.")

    cards(slide, [
        {"badge": "◆", "title": "Reviewable AI",
         "accent": AI,
         "body": ["Model, prompt version and source references are stored with "
                  "every AI output."]},
        {"badge": "◆", "title": "Human approval",
         "accent": BRAND,
         "body": ["Nothing reaches a funder without a person's sign-off — "
                  "sections stay editable throughout."]},
        {"badge": "◆", "title": "Numbers are not generated",
         "accent": CYAN,
         "body": ["Tables, charts and period-to-period changes are computed "
                  "from verified findings."]},
        {"badge": "◆", "title": "Never blocked by AI",
         "accent": SUCCESS,
         "body": ["If a draft call fails or a claim cannot be grounded, the "
                  "section falls back to deterministic text."]},
    ], y=y, h=2.45, cols=4, gap=0.18, title_size=11.5, body_size=10)

    quote(slide, "The risk in AI-assisted reporting is an unsupported claim. "
                 "Our answer is architectural: the model never invents a figure, "
                 "and every number in a report can be traced back to a source.",
          y=y + 2.63, h=0.92, size=12.5)

    strip(slide, "What the board should take away:",
          highlight="we are selling a governed drafting assistant, not an "
                    "autonomous author.",
          y=y + 3.71, h=0.48)


# --------------------------------------------------------------------------- #
# Slide 16 — Trust and security
# --------------------------------------------------------------------------- #

def s16_trust(prs):
    slide = bg(blank(prs))
    y = hero(
        slide, 16, "TRUST & SECURITY",
        "Built for sensitive programme and grant data",
        "DonorDesk protects participant, beneficiary, research and programme "
        "information through tenant isolation, controlled access, traceable AI "
        "outputs and immutable audit records.")

    cards(slide, [
        {"badge": "◆", "title": "Multi-tenant by design",
         "accent": BRAND,
         "body": ["Tenant isolation is enforced in the database, so no "
                  "organisation can ever see another's data."]},
        {"badge": "◆", "title": "Immutable audit log",
         "accent": BRAND_DK,
         "body": ["Every change is written to a chained, verifiable trail — "
                  "ready for internal compliance and donor review."]},
        {"badge": "◆", "title": "Reviewable AI",
         "accent": CYAN,
         "body": ["Model, prompt version and source references are stored with "
                  "every AI output, so nothing is written without a trace."]},
        {"badge": "◆", "title": "Least-privilege roles",
         "accent": BRAND_DP,
         "body": ["Project-level assignments keep access tight, from field "
                  "staff to approvers and administrators."]},
    ], y=y, h=2.24, cols=4, gap=0.18, title_size=11.5, body_size=10)

    panel(slide, ML, y + 2.42, CW, 1.60, "Also in place", [
        ("Encryption at rest and secure uploads", "item"),
        ("Confidential labels and PII warnings, with optional anonymisation of "
         "beneficiary lists", "item"),
        ("Version history and export history for the whole grant", "item"),
        ("Bring-your-own storage: field files can stay in the organisation's own "
         "Google Drive", "item"),
    ], accent=SUCCESS, body_size=10.5)

    strip(slide, "▲  In progress before scale: encrypted off-host backups with "
                 "restore tests, availability monitoring and alerting, and the "
                 "tenant-isolation step baked into every release.",
          y=y + 4.18, h=0.56, fill="FEF3C7", color="92400E", accent=AMBER,
          size=10)


# --------------------------------------------------------------------------- #
# Slide 17 — Why we win
# --------------------------------------------------------------------------- #

def s17_why_win(prs):
    slide = bg(blank(prs))
    y = hero(
        slide, 17, "WHY WE WIN",
        "Not another AI writer. Not another NGO ERP.",
        "Generic tools can generate text. Systems of record can store data. "
        "Neither can prove that a specific claim is supported by evidence the "
        "donor accepts.")

    cards(slide, [
        {"badge": "×", "title": "Generic AI assistants",
         "accent": FAINT,
         "body": ["Can write prose. No logframe, no evidence library, no "
                  "compliance state, no approval trail — and no way to show the "
                  "funder where a number came from."]},
        {"badge": "×", "title": "NGO ERP and MIS platforms",
         "accent": AMBER,
         "body": ["Record delivery, finance and procurement well, but reporting "
                  "still leaves the system and is reassembled by hand."]},
        {"badge": "×", "title": "File storage and folder trees",
         "accent": MUTED,
         "body": ["Store documents. No structure, no linking to indicators, no "
                  "readiness view, no audit trail."]},
    ], y=y, h=2.56, cols=3, gap=0.20, title_size=12, body_size=10.5)

    strip(slide, "Where DonorDesk sits:",
          highlight="delivery records → evidence → donor requirements → AI "
                    "drafting → human approval → export, in one auditable path.",
          y=y + 2.74, h=0.50)

    cards(slide, [
        {"badge": "01", "title": "Sector reporting packs",
         "accent": BRAND,
         "body": ["Nutrition, food security, WASH, protection, education."]},
        {"badge": "02", "title": "Donor template library",
         "accent": CYAN,
         "body": ["Every review that structures a funder's template becomes "
                  "reusable."]},
        {"badge": "03", "title": "The evidence graph",
         "accent": BRAND_DK,
         "body": ["Claims must reconcile to records — the part that is hard to "
                  "copy."]},
        {"badge": "04", "title": "Switching cost",
         "accent": AI,
         "body": ["Years of evidence, versions and audit history live inside "
                  "the workspace."]},
    ], y=y + 3.42, h=1.42, cols=4, gap=0.18, title_size=11.5, body_size=9.5)


# --------------------------------------------------------------------------- #
# Slide 18 — Status
# --------------------------------------------------------------------------- #

def s18_status(prs):
    slide = bg(blank(prs))
    y = hero(
        slide, 18, "STATUS", "Version 2.0 is live in production",
        "Deployed, verified and reversible today — with the remaining operational "
        "gaps listed here rather than discovered later.")

    lw = 5.52
    panel(slide, ML, y, lw, 3.86, "In production today", [
        ("donordesk.online served over HTTPS", "item"),
        ("API, web application and AI workers running as supervised services",
         "item"),
        ("Tenant isolation verified against the production database", "item"),
        ("Tiers, checkout and live billing provider in production", "item"),
        ("SuperAdmin control plane for tenants, plans and providers", "item"),
        ("One release path with a pre-deploy snapshot and rollback target",
         "item"),
    ], accent=SUCCESS, body_size=11)

    rx = ML + lw + 0.22
    rw = CW - lw - 0.22
    half = 1.86
    panel(slide, rx, y, rw, half, "Measured before every release", [
        ("25 / 25 report-quality acceptance cases pass", "item"),
        ("55 / 55 AI worker tests pass", "item"),
        ("136 / 137 infrastructure tests pass (1 intentionally skipped)",
         "item"),
        ("Full typecheck and build clean across all packages", "item"),
    ], accent=BRAND, body_size=10.5)

    panel(slide, rx, y + half + 0.14, rw, 3.86 - half - 0.14,
          "Operational gaps we are closing next", [
        ("Encrypted off-host backups with a documented restore test", "warn"),
        ("Availability monitoring with automatic restart and alerting", "warn"),
        ("Queue-backed background jobs for scale", "warn"),
        ("Transactional email for invitations and deadline reminders", "warn"),
    ], accent=AMBER, body_size=10.5)

    strip(slide, "Institutional position:",
          highlight="the product risk is behind us; what remains is operational "
                    "hardening, not invention.",
          y=y + 4.02, h=0.48)


# --------------------------------------------------------------------------- #
# Slide 19 — Roadmap
# --------------------------------------------------------------------------- #

def s19_roadmap(prs):
    slide = bg(blank(prs))
    y = hero(
        slide, 19, "ROADMAP", "Three horizons, each with an exit test",
        "What we are proving now, what makes it commercial, and what makes it the "
        "default choice for grant-funded reporting.")

    timeline(slide, [
        {
            "kicker": "NOW · 0–3 MONTHS", "title": "Prove it in the field",
            "accent": BRAND, "exit_fill": TINT,
            "bullets": [
                "Pilot with 3–5 NGOs, service-assisted at the start",
                "Enable AI-assisted drafting for pilot tenants",
                "Close the operational gaps: backups, monitoring, queue",
                "Onboarding smooth enough for a first report unaided",
                "Turn pilot feedback into the repeatable onboarding flow",
            ],
            "exit": "Three pilots each complete a live reporting period end to "
                    "end.",
        },
        {
            "kicker": "NEXT · 3–9 MONTHS", "title": "Make it commercial",
            "accent": CYAN, "exit_fill": TINT2,
            "bullets": [
                "Deadline reminders and escalation to owners",
                "Version history, comments and team collaboration",
                "Support console and tenant health for scale",
                "Self-serve growth on published pricing",
                "Reference customers for the fundraising story",
            ],
            "exit": "Paid conversion without service-assisted set-up.",
        },
        {
            "kicker": "LATER · 9–24 MONTHS", "title": "Own the category",
            "accent": AI, "dark": True, "exit_fill": "EDE9FE",
            "bullets": [
                "Sector packs: nutrition, food security, WASH, protection, "
                "education",
                "Integrations for field data, storage and dashboards",
                "Enterprise: SSO, data residency, service commitments",
                "A donor-side view of what was submitted, and when",
            ],
            "exit": "Teams choose DonorDesk because the pack already knows their "
                    "funder's requirements.",
        },
    ], y=y, h=4.86, gap=0.22)


# --------------------------------------------------------------------------- #
# Slide 20 — Business model and the ask
# --------------------------------------------------------------------------- #

def s20_business(prs):
    slide = bg(blank(prs))
    y = hero(
        slide, 20, "BUSINESS MODEL & DECISIONS",
        "Free to start. Priced for the teams that report most.",
        "Subscriptions scale with projects, seats, storage and AI drafts — "
        "supported by service-assisted onboarding that turns pilots into "
        "recurring revenue.")

    table(slide, ["Plan", "Monthly", "Annual", "Who it fits"], [
        ["Starter", "$0", "$0", "First funded programme"],
        ["Team", "$59", "$590", "Several grants or projects"],
        ["Growth", "$149", "$1,490", "Multiple funders, high volume"],
        ["Enterprise", "Custom", "Custom", "INGOs, research, public bodies"],
    ], x=ML, y=y + 0.06, w=6.44, size=10, head_size=10, row_h=0.40,
        col_widths=[1.5, 1.0, 1.0, 2.6])

    tf = textbox(slide, ML, y + 2.10, 6.44, 0.7)
    para(tf, "Annual billing gives two months free. Nonprofit discounts are "
             "available for qualifying organisations; tax is calculated at "
             "checkout.", size=9.5, color=MUTED, first=True, space_after=0,
         line_spacing=1.10)

    panel(slide, ML, y + 2.60, 6.44, 1.32, "Unit economics we watch", [
        ("Cost to serve falls as templates and sector packs are reused", "item"),
        ("AI usage is metered per plan, and a failed run is never charged",
         "item"),
        ("Storage stays bring-your-own where the organisation prefers it",
         "item"),
    ], accent=CYAN, body_size=10.5)

    rx = ML + 6.66
    rw = CW - 6.66
    half = 1.72
    panel(slide, rx, y + 0.06, rw, half, "How revenue builds", [
        ("Pilots sold as packages ($500–$2,000) that convert to subscriptions",
         "item"),
        ("Subscriptions scale by projects, seats, storage and AI-draft "
         "allowances", "item"),
        ("Services: template configuration, onboarding, training, review",
         "item"),
        ("A failed AI run is never billed — allowances match delivered value",
         "item"),
    ], accent=BRAND, body_size=10.5)

    panel(slide, rx, y + half + 0.18, rw, half, "Decisions requested today", [
        "1  Approve the published pricing and the nonprofit discount policy.",
        "2  Approve a 3–5 NGO pilot programme with service-assisted onboarding.",
        "3  Note the go-to-market sequence: local and national NGOs first.",
        "4  Note the 90-day operational hardening list on slide 18.",
    ], accent=AI, body_size=10.5)

    strip(slide, "Closing:",
          highlight="make every reporting period easier than the last.",
          y=y + 4.15, h=0.48)


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #

SLIDES = [s01_title, s02_exec_summary, s03_problem, s04_who, s05_vision,
          s06_how, s07_before_after, s08_templates, s09_drafting,
          s10_readiness, s11_evidence, s12_logframe, s13_exports, s14_memory,
          s15_responsible_ai, s16_trust, s17_why_win, s18_status,
          s19_roadmap, s20_business]


def main():
    parser = argparse.ArgumentParser(
        description="Build the DonorDesk board-approval deck.")
    parser.add_argument("--out",
                        default="DonorDesk-Board-Briefing-v2.0.pptx",
                        help="output .pptx path")
    args = parser.parse_args()

    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    for builder in SLIDES:
        builder(prs)
    prs.save(args.out)
    print(f"wrote {args.out} ({len(SLIDES)} slides)")


if __name__ == "__main__":
    main()

