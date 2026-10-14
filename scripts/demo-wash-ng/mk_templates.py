from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem
def mk(path, title, intro, sections):
    d = Document(); d.add_heading(title, 0); d.add_paragraph(intro)
    for lvl, h, ins, qs in sections:
        d.add_heading(h, lvl); d.add_paragraph(ins)
        for q in qs: d.add_paragraph(q, style="List Bullet")
    d.save(path)
mk("templates/LCBRF-monthly-flash-report-template.docx", "Lake Chad Basin Resilience Fund — Monthly Flash Report (Grant LCBRF/NG/2026/WASH-031)",
 "Grantees send a short monthly flash report to the Fund Management Unit within 7 days of month end. Report only verified results against the project logframe. Write in plain, factual English in the third person, refer to 'the project', and do not use promotional language. Disaggregate people-level results by sex.",
 [(1,"Headline Summary","Summarise the month in no more than 200 words: results, the main problem and the main decision needed.",["What was the most important result of the month?"]),
  (1,"Progress Against Logframe","Report progress against the logframe indicators.",["Provide a table: indicator, this month, cumulative, target, percent of target.","Explain any indicator behind plan."]),
  (1,"Activities Completed","List the activities completed by LGA with dates and participants.",["Which planned activities were not carried out, and why?"]),
  (1,"Water Quality and Public Health Risks","Report water quality results and any outbreak risks, including cholera alerts.",["Report the percentage of water points meeting the chlorine residual standard.","Which water quality results were not measured, and why?"]),
  (1,"Access, Security and Operating Environment","Describe access constraints, security incidents and how the project adapted.",["Any request for support or a change that needs the Fund's agreement?"]),
  (1,"Protection, Gender and Accountability to Affected Populations","Describe how protection, gender equality and accountability to affected populations were applied this month and report complaints received and resolved.",["Provide sex-disaggregated results.","Any safeguarding incident or protection concern?"]),
  (1,"Plan for Next Month","State the planned activities and expected results for the coming month.",[])])
mk("templates/LCBRF-final-report-template.docx", "Lake Chad Basin Resilience Fund — Final Narrative and Financial Report (Grant LCBRF/NG/2026/WASH-031)",
 "The final report covers the whole duration of the project and is submitted within 60 days of its end date. It assesses results against the logframe and the Fund's WASH standards (SPHERE minimum standards, national WASH cluster guidance). Use a factual, evidence-based register. Report life-of-project figures from verified records.",
 [(1,"Executive Summary","Summarise the project, its main results and overall performance in no more than 350 words.",[]),
  (1,"Context and Project Description","Describe the context and the design of the project.",[]),
  (2,"Context Update","Describe the situation in Jere, Konduga, Mafa and Dikwa LGAs and any significant change during implementation.",["Who were the people reached and how were they selected?"]),
  (2,"Theory of Change","Describe how outputs were expected to lead to outcomes and impact and whether the logic held.",[]),
  (1,"Results Achieved","Report life-of-project results against the logframe.",[]),
  (2,"Impact and Outcome Results","Report impact and outcome indicators.",["Provide an indicator table: baseline, target, achieved, percent of target.","Explain any target that was not met."]),
  (2,"Output Delivery","Report delivery of every output, disaggregated by sex where applicable.",["Provide a table of output indicators: target, achieved, percent of target."]),
  (2,"Deviations from Plan","Explain deviations from the work plan and any variance in results.",["What caused each deviation and what was done about it?"]),
  (1,"Compliance with WASH Standards","State how the water points, latrines and hygiene work met the SPHERE minimum standards and national guidance.",["Report water quality compliance over the project."]),
  (1,"Protection, Gender and Accountability","Report on protection, gender equality, inclusion and accountability to affected populations.",["Provide sex-disaggregated results and explain gender gaps.","Summarise complaints received and resolved."]),
  (1,"Sustainability and Handover","Explain what was handed over to the committees and the LGA water boards, and how services will be sustained.",["Which assets, systems and responsibilities were transferred?"]),
  (1,"Risks and Lessons Learned","Report which risks materialised and the main lessons for future programming.",["What would you do differently?"]),
  (1,"Financial Report","Provide a summary of the budget and expenditure of the project.",["Include a table: budget line, budget, expended, balance.","Explain any variance above ten percent."]),
  (1,"Recommendations","Recommend priorities for the Fund, the Borno State government and implementing partners.",[])])
# Quarterly report as a PDF with heading styles
ss = getSampleStyleSheet(); H1 = ParagraphStyle("h1", parent=ss["Heading1"], fontSize=18, spaceAfter=10); H2 = ParagraphStyle("h2", parent=ss["Heading2"], fontSize=14, spaceAfter=6); B = ss["BodyText"]
doc = SimpleDocTemplate("templates/LCBRF-quarterly-technical-report-template.pdf", pagesize=A4, title="LCBRF Quarterly Technical Report")
S = [Paragraph("Lake Chad Basin Resilience Fund — Quarterly Technical Report", H1), Paragraph("Grant LCBRF/NG/2026/WASH-031. Grantees submit a quarterly technical report within 15 days of the end of each quarter. Report verified results against the logframe in plain, factual English, in the third person. Do not use promotional language. Disaggregate people-level results by sex and state clearly which results were not measured.", B), Spacer(1, 10)]
secs = [(H1,"1. Executive Summary","Summarise the quarter in no more than 300 words: the main results, the main problems and the adaptation made.",[]),
 (H1,"2. Context and Needs Update","Describe changes in the humanitarian context, access and needs in the four LGAs during the quarter.",["Which new needs or risks arose?"]),
 (H1,"3. Results by Outcome","Report outcome indicators and explain the movement from the baseline.",[]),
 (H2,"3.1 Outcome 1: Safe water and sanitation","Report progress on water points, latrines and WASH blocks, with a results table: indicator, quarter value, cumulative, target, percent of target.",["Explain any output behind plan."]),
 (H2,"3.2 Outcome 2: Hygiene and management","Report hygiene promotion reach by sex, committee training and functionality.",[]),
 (H1,"4. Compliance with Sector Standards","State how the water points and sanitation facilities met SPHERE minimum standards and the national WASH cluster guidance.",["Report the water quality results and the percentage of water points that met the chlorine residual standard."]),
 (H1,"5. Protection, Gender and Accountability to Affected Populations","Report on protection, gender equality and inclusion, and summarise complaints received and resolved.",["Any safeguarding incident or protection concern?"]),
 (H1,"6. Financial Summary","Provide budget, expenditure and balance by budget line for the quarter, and explain any variance above ten percent.",[]),
 (H1,"7. Risks, Lessons and Adaptation","Update the risk register, state the main lessons and any change needing the Fund's agreement.",[]),
 (H1,"8. Plan for Next Quarter","State the planned activities and expected results.",[])]
for st, h, ins, qs in secs:
    S.append(Paragraph(h, st)); S.append(Paragraph(ins, B))
    if qs: S.append(ListFlowable([ListItem(Paragraph(q, B)) for q in qs], bulletType="bullet"))
    S.append(Spacer(1, 8))
doc.build(S); print("ok")
