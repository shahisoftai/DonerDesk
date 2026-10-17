#!/usr/bin/env python3
"""Build the management-first DonorDesk overview deck (v2.003).

Every factual line is taken from memorybank/docs/marketing/02-feature-inventory.md
(verified against the code) and the plan catalogue in
packages/domain/src/contexts/billing/plan.ts. Update those first, then this script.

Usage:
    pip3 install --target /tmp/pptxlib python-pptx
    PYTHONPATH=/tmp/pptxlib python3 scripts/build-donordesk-management-deck.py \
        --out memorybank/docs/presentations/DonorDesk-Overview-v2.003-Management-First.pptx
"""

from __future__ import annotations

import argparse

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

FONT = "Calibri"
INK = RGBColor(0x0F, 0x17, 0x2A)
BODY = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
BLUE = RGBColor(0x25, 0x63, 0xEB)
BLUE_DK = RGBColor(0x1E, 0x3A, 0x8A)
CYAN = RGBColor(0x06, 0xB6, 0xD4)
DARK = RGBColor(0x0B, 0x12, 0x20)
CARD = RGBColor(0xF1, 0xF5, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AMBER = RGBColor(0xB4, 0x53, 0x09)

SW, SH = Inches(13.333), Inches(7.5)


def _text(slide, x, y, w, h, text, size=18, bold=False, color=BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def _bullets(slide, x, y, w, h, items, size=18, color=BODY):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        run = p.add_run()
        run.text = f"•  {item}"
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = color
    return box


def _rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    shp = slide.shapes.add_shape(shape, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = 0.06
    shp.shadow.inherit = False
    return shp


def _footer(slide, n):
    _text(slide, Inches(0.6), Inches(7.0), Inches(9), Inches(0.3),
          "DonorDesk · Overview v2.003 · donordesk.online", size=10, color=MUTED)
    _text(slide, Inches(12.0), Inches(7.0), Inches(0.8), Inches(0.3), str(n), size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def _header(slide, kicker, title):
    _text(slide, Inches(0.6), Inches(0.45), Inches(12), Inches(0.35), kicker.upper(), size=12, bold=True, color=BLUE)
    _text(slide, Inches(0.6), Inches(0.8), Inches(12.1), Inches(1.0), title, size=30, bold=True, color=INK)


def _cards(slide, items, top=Inches(2.0), cols=3, card_h=Inches(2.2)):
    gap = Inches(0.3)
    w = int((SW - Inches(1.2) - gap * (cols - 1)) / cols)
    for i, (head, body) in enumerate(items):
        r, c = divmod(i, cols)
        x = Inches(0.6) + c * (w + gap)
        y = top + r * (card_h + gap)
        _rect(slide, x, y, w, card_h, CARD)
        _text(slide, x + Inches(0.25), y + Inches(0.2), w - Inches(0.5), Inches(0.5), head, size=17, bold=True, color=BLUE_DK)
        _text(slide, x + Inches(0.25), y + Inches(0.75), w - Inches(0.5), card_h - Inches(0.9), body, size=13, color=BODY)


def build(out_path: str) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH
    blank = prs.slide_layouts[6]
    n = 0

    def new_slide():
        nonlocal n
        n += 1
        return prs.slides.add_slide(blank)

    # 1. Title
    s = new_slide()
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = DARK
    _rect(s, Inches(0.6), Inches(1.6), Inches(0.12), Inches(3.2), CYAN, shape=MSO_SHAPE.RECTANGLE)
    _text(s, Inches(1.0), Inches(1.5), Inches(11), Inches(0.5), "DONORDESK · OVERVIEW", size=14, bold=True, color=CYAN)
    _text(s, Inches(1.0), Inches(2.1), Inches(11.5), Inches(1.6), "Run your programme.", size=54, bold=True, color=WHITE)
    _text(s, Inches(1.0), Inches(3.2), Inches(11.5), Inches(1.6), "Report from the record.", size=54, bold=True, color=CYAN)
    _text(s, Inches(1.0), Inches(4.7), Inches(11), Inches(0.9),
          "Programme management for grant-funded organisations, with donor reports drafted from the record and approved by people.",
          size=18, color=RGBColor(0xCB, 0xD5, 0xE1))
    _text(s, Inches(1.0), Inches(6.4), Inches(11), Inches(0.4), "October 2026 · donordesk.online", size=13, color=MUTED)

    # 2. The problem
    s = new_slide()
    _header(s, "The problem", "A programme is scattered, then rebuilt for every donor report")
    _bullets(s, Inches(0.6), Inches(2.0), Inches(12), Inches(4.5), [
        "Indicators live in one spreadsheet, activities in another, evidence in folders and chat threads.",
        "The report is assembled in the last weeks before the deadline, from numbers nobody can trace back.",
        "Compliance gaps (a missing document, an unverified figure) surface on the day of submission.",
        "Each new donor wants the same facts in a different template, so the work is repeated.",
    ], size=20)
    _footer(s, n)

    # 3. The record
    s = new_slide()
    _header(s, "The shift", "One record from plan to report")
    _cards(s, [
        ("1 · Plan", "Logframe, indicators, baselines and targets. Import what you already have from Excel, CSV or Google Sheets."),
        ("2 · Record", "Field staff log activities as they happen, typed or dictated. Indicator values are entered once per period."),
        ("3 · Check", "Evidence is linked to the activities it proves. The checklist shows what is missing and what closes it."),
        ("4 · Report", "The donor template structures the report. Drafts cite their sources. A person approves before export."),
    ], cols=4, card_h=Inches(3.0))
    _text(s, Inches(0.6), Inches(6.2), Inches(12), Inches(0.5),
          "The record is the working system. The report is what it produces.", size=16, bold=True, color=BLUE_DK)
    _footer(s, n)

    # 4. Run the programme
    s = new_slide()
    _header(s, "Run the programme", "The working system your team uses every day")
    _cards(s, [
        ("Logframe and indicators", "Goals to activities, with indicators, baselines, targets and frequency. Archive rather than delete once values exist."),
        ("Activity capture", "Date, place, participants and a narrative. The reporting period is worked out from the date. Dictation and offline saving."),
        ("Data Center", "Every indicator value across every period in one sheet. Download, change, import with a preview. Charts from the data."),
        ("Evidence library", "Files in your Google Drive or managed storage. Linked to activities and indicators. Evidence packs as ZIP."),
        ("Team and roles", "Seven roles: Admin, Project Manager, M&E, Grants, Field, Compliance and Viewer. Seats follow the plan."),
        ("My Work", "One queue across every project: reports to review, checklist items, evidence to verify, activities to accept."),
    ], cols=3, card_h=Inches(2.2))
    _footer(s, n)

    # 5. Activity capture
    s = new_slide()
    _header(s, "Field capture", "Activities recorded when the work happens")
    _bullets(s, Inches(0.6), Inches(2.0), Inches(6.3), Inches(4.6), [
        "Date, place, participants with breakdowns, and a narrative. A record can cover a date span.",
        "Dictate the narrative with the microphone button. The form fills in date, place and people only from what the words state.",
        "A finished record is saved on the device and sent when the connection returns.",
        "Managers accept, send back or reject records, one at a time or in bulk.",
    ], size=17)
    _rect(s, Inches(7.4), Inches(2.0), Inches(5.3), Inches(4.4), CARD)
    _text(s, Inches(7.7), Inches(2.2), Inches(4.8), Inches(0.5), "Status workflow", size=17, bold=True, color=BLUE_DK)
    _text(s, Inches(7.7), Inches(2.8), Inches(4.8), Inches(3.4), [
        "Draft → Submitted → Accepted",
        "Submitted → Needs revision → Submitted",
        "Rejected or Withdrawn from most states",
        "",
        "Undoing an accepted record is refused when an approved report already used it.",
    ], size=14, color=BODY)
    _text(s, Inches(0.6), Inches(6.5), Inches(12), Inches(0.4),
          "Limit: the app needs a connection to open. Offline use is queued saving only.", size=12, color=AMBER)
    _footer(s, n)

    # 6. Data Center
    s = new_slide()
    _header(s, "Data Center", "Every indicator value, every period, one sheet")
    _bullets(s, Inches(0.6), Inches(2.0), Inches(12), Inches(4.5), [
        "See all reporting periods side by side, by period or as a list.",
        "Download the sheet to Excel or CSV. A missing value stays blank, never zero.",
        "Import the edited sheet and see every change before it applies. Empty cells never clear a stored value.",
        "A problem in any row blocks the whole import, so a half-applied update cannot happen.",
        "Draw charts from the data. A chart added to a report is a snapshot you update when you choose.",
    ], size=19)
    _text(s, Inches(0.6), Inches(6.5), Inches(12), Inches(0.4),
          "Finance figures appear only when finance mode is on for the project (off by default).", size=12, color=AMBER)
    _footer(s, n)

    # 7. Evidence
    s = new_slide()
    _header(s, "Evidence", "Evidence that stays linked to the work it proves")
    _cards(s, [
        ("Where files live", "Your own Google Drive (linked files do not count toward storage), or DonorDesk managed storage within your plan."),
        ("Linked to the record", "Attach a file to an activity or an indicator value, and mark it verified. One linker does this everywhere."),
        ("Read by the drafts", "Text is read from Word, PDF, Excel, CSV, text and Markdown. Cited text is fingerprinted, so a later change to the source is detectable."),
    ], cols=3, card_h=Inches(2.6))
    _text(s, Inches(0.6), Inches(5.3), Inches(12), Inches(1.2), [
        "Evidence packs export as a ZIP with a README. Files that cannot be included (Drive-only or unreadable) are named in the README.",
        "Every change to the record is kept in a hash-chained audit log.",
    ], size=15, color=BODY)
    _footer(s, n)

    # 8. Compliance
    s = new_slide()
    _header(s, "Compliance", "Gaps found weeks before the deadline, not on the day")
    _cards(s, [
        ("Live checklist", "Every requirement from the donor template and the project's rules, with the blocker and the action that closes it."),
        ("Closes by itself", "A verified procurement document closes its item. Items that need a decision stay open until a person decides."),
        ("Readiness by stage", "Drafting, review and submission. A confidentiality check blocks submission while it is open."),
    ], cols=3, card_h=Inches(2.6))
    _text(s, Inches(0.6), Inches(5.3), Inches(12), Inches(1.0),
          "Readiness guides what remains. It does not replace the donor's own review.", size=15, color=BODY)
    _footer(s, n)

    # 9. Reports
    s = new_slide()
    _header(s, "Donor reports", "Drafted from the record, approved, then exported")
    _cards(s, [
        ("Donor templates", "Upload Word, PDF, text, Markdown, Excel or CSV. Sections, instructions, required tables and page limits are extracted for your review."),
        ("Sourced drafting", "Each section follows the approved template and cites its sources. A figure not in the inputs is rejected and regenerated, or the section falls back and says why."),
        ("Approve and export", "Second-approver rules, a sealed submission snapshot, exports to Word, PDF, Excel and ZIP. A watermarked internal copy for review."),
    ], cols=3, card_h=Inches(2.6))
    _text(s, Inches(0.6), Inches(5.3), Inches(12), Inches(1.2), [
        "Period types: monthly, quarterly, semi-annual, annual, final and custom. Roll-up reports show life-of-project totals and percent of target.",
        "An approved report can be reopened, with a reason, and approved again.",
    ], size=15, color=BODY)
    _footer(s, n)

    # 10. Trust
    s = new_slide()
    _header(s, "Trust", "People decide. The record checks the numbers.")
    _cards(s, [
        ("Verified numbers", "Tables and charts come from verified data. The AI writes prose only, and every figure is checked against the inputs."),
        ("Traceable AI", "Each AI output stores the model, prompt version and sources. Failed or timed-out drafts are not billed."),
        ("Isolation and roles", "Row-level security keeps each organisation separate. Least-privilege roles per project. Hash-chained audit log."),
    ], cols=3, card_h=Inches(2.6))
    _text(s, Inches(0.6), Inches(5.3), Inches(12), Inches(1.2), [
        "Published: data processing agreement, subprocessor list, privacy and security pages.",
        "No external certifications are held. None are claimed.",
    ], size=15, color=BODY)
    _footer(s, n)

    # 11. Who
    s = new_slide()
    _header(s, "Who it's for", "Built for grant-funded programmes and their reporting teams")
    _cards(s, [
        ("M&E Officer", "Enter each value once; every report and chart reads the same data."),
        ("Grants Officer", "See readiness and blockers weeks before the deadline."),
        ("Project Manager", "Review and accept field records in one place."),
        ("Field Officer", "Record what happened in a few lines, or by dictation."),
        ("Compliance Officer", "Every claim links to its evidence; the audit log shows who changed what."),
        ("Approver", "Second-approver rules; the submitted file is sealed."),
    ], cols=3, card_h=Inches(2.0))
    _footer(s, n)

    # 12. Pricing
    s = new_slide()
    _header(s, "Plans", "Start free. Upgrade when the programme grows.")
    rows = [
        ["", "Starter", "Team", "Growth", "Enterprise"],
        ["Monthly", "Free", "$79", "$149", "Contract"],
        ["Annual (two months free)", "Free", "$790", "$1,490", "From $12,000 a year"],
        ["NGO price, first year (25% off)", "—", "$59.25 / month", "$111.75 / month", "In the contract"],
        ["Active projects", "1", "5", "20", "Contracted"],
        ["Full seats", "1", "5", "15", "Contracted"],
        ["Viewer seats", "2", "Unlimited", "Unlimited", "Unlimited"],
        ["Managed storage", "1 GB", "25 GB", "100 GB", "Contracted"],
        ["AI report drafts / month", "5", "20", "100", "Contracted pool"],
        ["Own AI provider", "—", "—", "Yes", "Yes"],
        ["SSO / SCIM", "—", "—", "—", "Yes"],
    ]
    tbl = s.shapes.add_table(len(rows), 5, Inches(0.6), Inches(1.85), Inches(12.1), Inches(4.3)).table
    widths = [3.6, 2.0, 2.2, 2.2, 2.1]
    for i, wdt in enumerate(widths):
        tbl.columns[i].width = Inches(wdt)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = val
            para = cell.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
            for run in para.runs:
                run.font.name = FONT
                run.font.size = Pt(13)
                run.font.bold = r == 0 or c == 0
                run.font.color.rgb = WHITE if r == 0 else BODY
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLUE_DK if r == 0 else (CARD if r % 2 else WHITE)
    _text(s, Inches(0.6), Inches(6.62), Inches(12), Inches(0.35),
          "Prices from the billing catalogue. Enterprise is by contract. Trial availability is set per launch.",
          size=11, color=MUTED)
    _footer(s, n)

    # 13. Limits
    s = new_slide()
    _header(s, "Honest limits", "What DonorDesk does not do, stated up front")
    _bullets(s, Inches(0.6), Inches(2.0), Inches(12), Inches(4.8), [
        "It does not send reports to donors. You upload the approved file.",
        "The app does not open offline. Finished activities are queued and sent later.",
        "Dictation uses the browser's speech recognition (Chrome, Edge, Safari).",
        "Donor-template Word exports do not include charts. Charts appear in the standard exports.",
        "Finance figures are off by default and appear only when switched on for a project.",
        "Enterprise SSO (OIDC) and SCIM are set up with support, not self-serve.",
        "No customer case studies or certifications are claimed in this deck.",
    ], size=18)
    _footer(s, n)

    # 14. Next steps
    s = new_slide()
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = DARK
    _text(s, Inches(1.0), Inches(1.6), Inches(11), Inches(0.5), "NEXT STEP", size=14, bold=True, color=CYAN)
    _text(s, Inches(1.0), Inches(2.2), Inches(11.5), Inches(1.4), "See it on a real programme.", size=48, bold=True, color=WHITE)
    _bullets(s, Inches(1.0), Inches(3.9), Inches(11), Inches(2.4), [
        "Take the narrated product tour: donordesk.online/tour",
        "Create a free workspace; the demo project is included and uses no AI credits: donordesk.online/signup",
        "Read every feature with its limits: donordesk.online/features",
        "Enterprise and pilots: donordesk.online/contact-sales",
    ], size=18, color=RGBColor(0xCB, 0xD5, 0xE1))

    prs.save(out_path)
    print(f"wrote {out_path} ({n} slides)")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", required=True)
    build(parser.parse_args().out)
